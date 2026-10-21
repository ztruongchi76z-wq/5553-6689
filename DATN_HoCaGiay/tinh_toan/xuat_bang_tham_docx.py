# -*- coding: utf-8 -*-
"""
XUẤT BẢNG KẾT QUẢ TÍNH THẤM SAU SỬA CHỮA RA WORD – MẶT CẮT LÒNG SÔNG VÀ MẶT CẮT SƯỜN ĐỒI
Cùng mẫu Bảng 5.1 của thuyết minh (Times New Roman, chữ trong bảng 10 pt, hàng tiêu đề nền xám nhạt, tổng bề rộng
như bảng gốc) để SV dán thẳng vào bản Word.
Số liệu: ket_qua/tham.json (chay_tham_on_dinh.py – lòng sông), ket_qua/tham_suon_doi.json (tham_suon_doi.py –
sườn đồi, mặt cắt V – T9). Chạy lại hai chương trình đó trước nếu thông số đầu vào thay đổi.
  python3 xuat_bang_tham_docx.py
"""
import json
import os
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import so_lieu as S

DIR = os.path.dirname(os.path.abspath(__file__))
RA = os.path.join(DIR, '..', 'TAI_VE', 'BangTham_LongSong_SuonDoi.docx')
COT = [1836, 1180, 1180, 1336, 1324, 1113, 1102]     # bề rộng cột (dxa), tổng như Bảng 5.1 trong thuyết minh;
                                                     # cột J cửa ra rộng hơn để 0,017 / 0,051 không xuống dòng
FONT = 'Times New Roman'


def so(x, n=2):
    return ('%.*f' % (n, x)).replace('.', ',').replace('-', '−')


def so_q(q):
    """1,67·10^{-6} – cùng cách viết với Bảng 5.1 (số mũ chỉ số trên)."""
    return '%s·10^{-6}' % so(q * 1e6)


def chu(p, markup, co=13, dam=False, nghieng=False):
    """Thêm chữ vào đoạn; ^{...} là chỉ số trên."""
    for i, phan in enumerate(re.split(r'\^\{([^}]*)\}', markup)):
        if not phan:
            continue
        r = p.add_run(phan)
        r.font.name, r.font.size, r.bold, r.italic = FONT, Pt(co), dam, nghieng
        r._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        r.font.superscript = bool(i % 2)
    return p


def doan(d, markup='', co=13, can=WD_ALIGN_PARAGRAPH.JUSTIFY, dam=False, nghieng=False, truoc=6, sau=0,
         thut=0.0):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.alignment, pf.space_before, pf.space_after, pf.line_spacing = can, Pt(truoc), Pt(sau), 1.2
    pf.first_line_indent = Cm(thut) if thut else None
    return chu(p, markup, co, dam, nghieng)


def o_bang(cell, markup, rong, dam=False, can=WD_ALIGN_PARAGRAPH.CENTER, nen=None):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn('w:tcW'))
    if tcw is None:
        tcw = OxmlElement('w:tcW')
        tcpr.append(tcw)
    tcw.set(qn('w:w'), str(rong))
    tcw.set(qn('w:type'), 'dxa')
    if nen:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), nen)
        tcpr.append(shd)
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.alignment, pf.space_before, pf.space_after, pf.line_spacing = can, Pt(1), Pt(1), 1.0
    chu(p, markup, 10, dam)


def hang_lap(row, tieu_de=False):
    trpr = row._tr.get_or_add_trPr()
    for tag in (['w:tblHeader'] if tieu_de else []) + ['w:cantSplit']:
        e = OxmlElement(tag)
        e.set(qn('w:val'), 'true')
        trpr.append(e)


def bang(d, chu_thich, tieu_de, du_lieu):
    p = doan(d, chu_thich, co=12, can=WD_ALIGN_PARAGRAPH.CENTER, nghieng=True, truoc=6, sau=6)
    p.paragraph_format.keep_with_next = True
    t = d.add_table(rows=1 + len(du_lieu), cols=len(COT))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), COT):
        gc.set(qn('w:w'), str(w))
    for j, (c, w) in enumerate(zip(t.rows[0].cells, COT)):
        o_bang(c, tieu_de[j], w, dam=True, nen='F2F2F2')
    hang_lap(t.rows[0], tieu_de=True)
    for i, dong in enumerate(du_lieu):
        for j, (c, w) in enumerate(zip(t.rows[i + 1].cells, COT)):
            o_bang(c, dong[j], w, can=WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
        hang_lap(t.rows[i + 1])
    return t


def tb95(r, n=2):
    return '%s / %s' % (so(r['tb'], n), so(r['p95'], n))


def main():
    T = json.load(open(os.path.join(DIR, 'ket_qua', 'tham.json'), encoding='utf-8'))['tham']
    SD = json.load(open(os.path.join(DIR, 'ket_qua', 'tham_suon_doi.json'), encoding='utf-8'))
    TH = [('MNDBT +%s' % so(S.MNDBT), 'MNDBT', 'TH1_MNDBT'), ('MNLTK +%s' % so(S.MNLTK), 'MNLTK', 'TH2_MNLTK'),
          ('MNLKT +%s' % so(S.MNLKT), 'MNLKT', 'TH3_MNLKT'), ('MNDBT, thoát nước bị tắc', 'MNDBT_tac', 'TH4_MNDBT_tac')]

    # ---- mặt cắt lòng sông
    ls = []
    for nhan, k, _ in TH:
        ht, sc = T['HT_' + k], T['SC_' + k]
        ls.append([nhan, so_q(ht['q']), so_q(sc['q']), '%s / %s' % (so(sc['J_tiep']), so(sc['J_tiep95'])),
                   '%s / %s' % (so(sc['J_ra'], 3), so(sc['J_ra95'], 3)), '%s / %s' % (so(sc['J_ck']), so(sc['J_ck95'])),
                   so(sc['z_tim'])])
    giam = [100 * (1 - T['SC_' + k]['q'] / T['HT_' + k]['q']) for _, k, _ in TH]
    ls_tiep = max(T['SC_' + k]['J_tiep95'] for _, k, _ in TH)
    ls_ra = max(T['SC_' + k]['J_ra95'] for _, k, _ in TH)

    # ---- mặt cắt sườn đồi (V – T9)
    sd = []
    for nhan, _, k in TH:
        ht, sc = SD[k], SD['SC_' + k]
        sd.append([nhan, so_q(ht['q']), so_q(sc['q']), tb95(sc['J_tiep']), tb95(sc['J_ra'], 3), tb95(sc['J_ck']),
                   so(sc['z_tim'])])
    sd_tiep = max(SD['SC_' + k]['J_tiep']['p95'] for _, _, k in TH)
    sd_ra = max(SD['SC_' + k]['J_ra']['p95'] for _, _, k in TH)
    sd_ck = max(SD['SC_' + k]['J_ck']['p95'] for _, _, k in TH)
    lech_q = max(abs(SD['SC_' + k]['q'] / SD[k]['q'] - 1) for _, _, k in TH) * 100
    ty_so = [T['SC_' + a]['q'] / SD['SC_' + b]['q'] for _, a, b in TH]
    LUOI = {Zw: [v for k, v in SD.items() if k.startswith('luoi_') and k.endswith('_%.2f' % Zw)] for Zw in (S.MNDBT, S.MNLKT)}
    nut = sorted(r['so_nut'] for r in LUOI[S.MNDBT])
    ht_q = max(max(r['q'] for r in v) / min(r['q'] for r in v) - 1 for v in LUOI.values()) * 100
    ht_z = max(max(r['z_tim'] for r in v) - min(r['z_tim'] for r in v) for v in LUOI.values())
    ht_tiep = [r['J_tiep']['p95'] for r in LUOI[S.MNLKT]]
    ht_ra = max(r['J_ra']['p95'] for v in LUOI.values() for r in v)

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.0), Cm(2.5), Cm(2.5)
    doan(d, 'KẾT QUẢ TÍNH THẤM SAU SỬA CHỮA – MẶT CẮT LÒNG SÔNG VÀ MẶT CẮT SƯỜN ĐỒI', co=13, dam=True,
         can=WD_ALIGN_PARAGRAPH.CENTER, truoc=0)
    doan(d, 'Thông số đầu vào thống nhất: đỉnh đập trước nâng cấp +%s m, sau sửa chữa +%s m; cơ thượng lưu, hạ lưu '
            '+%s m; MNDBT +%s m, MNLTK +%s m, MNLKT +%s m.'
         % (so(S.Z_DINH_DAP_HT), so(S.Z_DINH_DAP_SC), so(S.Z_CO_TL, 1), so(S.MNDBT), so(S.MNLTK), so(S.MNLKT)),
         co=11, nghieng=True, can=WD_ALIGN_PARAGRAPH.CENTER)

    tieu_de = ['Trường hợp', 'q hiện trạng (m^{3}/s/m)', 'q sau SC (m^{3}/s/m)', 'J tiếp xúc ống khói TB / 95%',
               'J cửa ra TB / 95%', 'J chân khay TB / 95%', 'Z bão hòa tại tim (m)']
    bang(d, 'Bảng 5.1 Kết quả tính thấm mặt cắt lòng sông sau sửa chữa', tieu_de, ls)
    doan(d, 'So với hiện trạng, lưu lượng thấm giảm %s ÷ %s%%; gradien thấm ở mọi bộ phận nhỏ hơn giá trị cho phép '
            '(tiếp xúc ống khói ≤ %s < 2,0; cửa ra ≤ %s < 0,85).'
         % (so(min(giam), 1), so(max(giam), 1), so(ls_tiep), so(ls_ra)), truoc=6, thut=1.0)

    tieu_de[3] = 'J tiếp xúc đệm tiêu nước TB / 95%'
    bang(d, 'Bảng 5.2 Kết quả tính thấm mặt cắt sườn đồi sau sửa chữa', tieu_de, sd)
    doan(d, 'Mặt cắt sườn đồi (mặt cắt V – T9) tính theo cùng phương pháp và điều kiện biên như mặt cắt lòng sông: '
            'khối lõi ĐĐ-I có chân khay đáy +63,8 m, khối thượng lưu và hạ lưu ĐĐ-II, đệm tiêu nước nằm ngang ở chân '
            'mái hạ lưu, nền sét pha tàn tích (lớp 4, k = 4,11·10^{-5} cm/s); hệ số thấm của lõi lấy theo thí nghiệm '
            'trong phòng k = 1·10^{-4} cm/s – lưu lượng thấm tính được sai khác −7 ÷ +13%% so với kết quả tính thấm '
            'mặt cắt V – T9 trong hồ sơ BCNCKT (Phụ lục 3). Mặt đập chỉ được nâng trong phạm vi bề rộng đỉnh '
            '(+%s → +%s m), phần đắp nâng nằm sau tường chắn sóng và cao hơn đường bão hòa nên dòng thấm sau sửa chữa '
            'hầu như không đổi so với hiện trạng (lưu lượng sai khác dưới %s%%). Lưu lượng thấm đơn vị nhỏ hơn mặt cắt lòng sông %s ÷ %s lần vì cột '
            'nước thấm nhỏ hơn (mặt đất hạ lưu +68,3 m, ở lòng sông +54,0 m). Gradien thấm nhỏ hơn giá trị cho phép: '
            'cửa ra ≤ %s < 0,85 (TCVN 8216:2018, Bảng 6, đất á sét, đập cấp II); tiếp xúc đệm tiêu nước ≤ %s < 2,0 '
            '(cùng giới hạn với tiếp xúc ống khói); chân khay ≤ %s.'
         % (so(S.Z_DINH_DAP_HT), so(S.Z_DINH_DAP_SC), so(max(lech_q, 0.1), 1), so(min(ty_so), 1), so(max(ty_so), 1),
            so(sd_ra), so(sd_tiep), so(sd_ck)), truoc=6, thut=1.0)
    doan(d, 'Kiểm tra chia mịn lưới (%s ÷ %s nút): lưu lượng thấm thay đổi không quá %s%%, Z bão hòa tại tim không '
            'quá %s m; phân vị 95%% của J tiếp xúc đệm tiêu nước khi MNLKT dao động %s ÷ %s (< 2,0), J cửa ra ≤ %s. '
            'Giá trị lớn nhất cục bộ tăng dần khi chia mịn lưới (tập trung ở góc đầu thượng lưu của đệm tiêu nước – '
            'điểm kỳ dị của lời giải) nên không dùng để đánh giá.'
         % (format(nut[0], ',').replace(',', '.'), format(nut[-1], ',').replace(',', '.'), so(ht_q, 1), so(ht_z),
            so(min(ht_tiep)), so(max(ht_tiep)), so(ht_ra)), truoc=6, thut=1.0)
    doan(d, 'Ghi chú: TB – trung bình theo diện tích các phần tử trong vùng bão hòa; 95% – phân vị 95%. Các cột '
            'gradien và Z bão hòa tại tim là giá trị sau sửa chữa. J tiếp xúc – phần tử đất kề ống khói (lòng sông) '
            'hoặc kề đệm tiêu nước (sườn đồi); J cửa ra – phần tử đất kề mặt thấm ra; J chân khay – phần tử lõi nằm '
            'trong chân khay. Hàng “thoát nước bị tắc”: thiết bị thoát nước có hệ số thấm như đất đắp ĐĐ-II.',
         co=11, nghieng=True, truoc=6)
    d.save(RA)
    print('Đã ghi', os.path.normpath(RA))
    for dong in ls + sd:
        print(' | '.join(dong))


if __name__ == '__main__':
    main()
