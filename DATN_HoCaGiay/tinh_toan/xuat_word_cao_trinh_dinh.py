# -*- coding: utf-8 -*-
"""
XUẤT WORD MỤC 2.3 – XÁC ĐỊNH CAO TRÌNH ĐỈNH ĐẬP (đánh giá an toàn phòng lũ) ĐỂ SV DÁN VÀO THUYẾT MINH
Mực nước theo yêu cầu SV 27/10/2026: MNDBT +74,70; MNLTK +77,43; MNLKT +78,39 (điều tiết lũ QTVH, vận hành bình
thường). Số liệu tính bằng cao_trinh_dinh.py – cùng công thức với file Excel xuat_excel_cao_trinh_dinh.py.
Công thức là phương trình Word (OMML), bảng Times New Roman 10 pt như các bảng đã gửi.
  python3 xuat_word_cao_trinh_dinh.py [đường dẫn .docx ra]
Mặc định: ../TAI_VE/Muc2.3_CaoTrinhDinhDap.docx
"""
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

import cao_trinh_dinh as CT
import so_lieu as S
from xuat_bang_tham_docx import bang, chu, doan, so

DIR = os.path.dirname(os.path.abspath(__file__))
RA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DIR, '..', 'TAI_VE', 'Muc2.3_CaoTrinhDinhDap.docx')
MNDBT, MNLTK, MNLKT = 74.70, 77.43, 78.39
A1, A2, A3 = CT.A_AN_TOAN['II']
D1, D2 = 2970.0, 3025.0
PT_MAI = 6.2                                   # phương vị pháp tuyến mái thượng lưu (°)
Z_DINH_HT, Z_DINH_KS, Z_TUONG, Z_DINH_SC = 78.50, (78.32, 78.44), 79.40, S.Z_DINH_DAP_SC
HUONG = [('N', 'Bắc', 0), ('NE', 'Đông Bắc', 45), ('NW', 'Tây Bắc', 315)]
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


# ------------------------------------------------------------------ tính toán
def goc_toi(az):
    return abs((az - PT_MAI + 180) % 360 - 180)


KQ = {}
for ma, ten, az in HUONG:
    k = S.GIO_HUONG.index(ma)
    a = goc_toi(az)
    r1 = CT.tinh_truong_hop(MNDBT, S.GIO[4][k], D1, alpha=a)
    r2 = CT.tinh_truong_hop(MNLTK, S.GIO[50][k], D2, alpha=a)
    KQ[ma] = dict(ten=ten, a=a, r1=r1, r2=r2, Z1=MNDBT + r1['dh'] + r1['hsl'] + A1,
                  Z2=MNLTK + r2['dh'] + r2['hsl'] + A2, Z3=MNLKT + A3)
    KQ[ma]['Zd'] = max(KQ[ma]['Z1'], KQ[ma]['Z2'], KQ[ma]['Z3'])
MA_MAX = max(KQ, key=lambda m: KQ[m]['Zd'])
ZD = KQ[MA_MAX]['Zd']
Z_LKT03 = MNLKT + 0.30
assert KQ[MA_MAX]['Z2'] == ZD                  # đoạn kết luận viết cho trường hợp Z₂ quyết định


# ------------------------------------------------------------------ phương trình Word (OMML)
def _m(tag, *con, **thuoc):
    e = etree.Element('{%s}%s' % (M, tag))
    for k, v in thuoc.items():
        e.set('{%s}%s' % (M, k), v)
    for c in con:
        e.append(c)
    return e


def t(chu_, dung=False):
    """m:r – dung=True: chữ đứng (tên mực nước, cos, số mũ chữ)."""
    r = _m('r')
    if dung:
        r.append(_m('rPr', _m('sty', val='p')))
    rpr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), 'Cambria Math')
    rf.set(qn('w:hAnsi'), 'Cambria Math')
    rpr.append(rf)
    r.append(rpr)
    tt = _m('t')
    tt.text = chu_
    tt.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(tt)
    return r


def _ds(x):
    return x if isinstance(x, list) else [x]


def sub(goc, cs):
    return _m('sSub', _m('e', *_ds(goc)), _m('sub', *_ds(cs)))


def sup(goc, mu):
    return _m('sSup', _m('e', *_ds(goc)), _m('sup', *_ds(mu)))


def phan_so(tu, mau):
    return _m('f', _m('num', *_ds(tu)), _m('den', *_ds(mau)))


def can(x):
    return _m('rad', _m('radPr', _m('degHide', val='1')), _m('deg'), _m('e', *_ds(x)))


def ngoac(x, mo='(', dong=')'):
    return _m('d', _m('dPr', _m('begChr', val=mo), _m('endChr', val=dong)), _m('e', *_ds(x)))


def gach(x):
    return _m('bar', _m('barPr', _m('pos', val='top')), _m('e', *_ds(x)))


def k_(cs):
    return sub(t('k'), t(cs))


def phuong_trinh(d, phan, so_ct):
    p = d.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(4), Pt(4), 1.2
    ppr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    for vt, kieu in ((4535, 'center'), (9071, 'right')):
        tb = OxmlElement('w:tab')
        tb.set(qn('w:val'), kieu)
        tb.set(qn('w:pos'), str(vt))
        tabs.append(tb)
    ppr.append(tabs)
    p.add_run('\t')
    p._p.append(_m('oMath', *phan))
    chu(p, '\t(%s)' % so_ct)
    return p


# ------------------------------------------------------------------ tài liệu
def tieu_de(d, chu_):
    return doan(d, chu_, dam=True, nghieng=True, truoc=12, sau=6, can=WD_ALIGN_PARAGRAPH.LEFT)


def main():
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.0), Cm(2.5), Cm(2.5)
    st = d.styles['Normal']
    st.font.name, st.font.size = 'Times New Roman', Pt(13)

    tieu_de(d, '2.3. Đánh giá an toàn phòng lũ')
    tieu_de(d, '2.3.1 Công thức tính toán')
    doan(d, 'Theo TCVN 8216:2018, cao trình đỉnh đập lấy giá trị lớn nhất trong ba trường hợp:', thut=1.0)
    phuong_trinh(d, [sub(t('Z'), t('1')), t(' = '), t('MNDBT', True), t(' + Δh + '), sub(t('h'), t('sl')), t(' + a')], '2-5')
    phuong_trinh(d, [sub(t('Z'), t('2')), t(' = '), t('MNLTK', True), t(' + Δh′ + '), sub(t('h′'), t('sl')), t(' + a′')],
                 '2-6')
    phuong_trinh(d, [sub(t('Z'), t('3')), t(' = '), t('MNLKT', True), t(' + a″')], '2-7')
    doan(d, 'Trong đó: Δh, Δh′ – độ dềnh do gió (m); h_{sl}, h′_{sl} – chiều cao sóng leo có mức bảo đảm 1%% (m), ứng với '
            'vận tốc gió tính toán lớn nhất có tần suất P = 4%% khi hồ ở MNDBT và vận tốc gió bình quân lớn nhất P = 50%% '
            'khi hồ ở MNLTK (TCVN 8216:2018, Bảng 4 – công trình cấp II); khi hồ ở MNLKT không xét sóng do gió; a, a′, a″ – '
            'độ vượt cao an toàn, theo TCVN 8216:2018 Bảng 3 với đập cấp II: a = %s m; a′ = %s m; a″ = %s m.'
            % (so(A1, 1), so(A2, 1), so(A3, 1)), thut=1.0)
    doan(d, 'Độ dềnh do gió và sóng leo được tính theo TCVN 8421:2010:', thut=1.0)
    phuong_trinh(d, [t('Δh = '), k_('w'), t('·'), phan_so([sup(t('V'), t('2')), t('·D')], t('g·H')), t('·'),
                     t('cos', True), t(' α ;     '), k_('w'), t(' = 3·'), sup(t('10'), t('−7')), t('·(1 + 0,3·V)')], '2-8')
    gh = [t('g·'), gach(t('h'))]
    phuong_trinh(d, [phan_so(gh, sup(t('V'), t('2'))), t(' = 0,16·'),
                     ngoac([t('1 − '), sup(ngoac([t('1 + 0,006·'), can(phan_so(t('gD'), sup(t('V'), t('2'))))]), t('−2'))],
                           '[', ']'),
                     t(' ;     '), phan_so([t('g·'), gach(t('τ'))], t('V')), t(' = 19,1·'),
                     sup(ngoac(phan_so([t('g·'), gach(t('h'))], sup(t('V'), t('2')))), t('0,625'))], '2-9')
    phuong_trinh(d, [sub(t('h'), t('sl1%')), t(' = '), k_('r'), t('·'), k_('p'), t('·'), k_('sp'), t('·'), k_('run'), t('·'),
                     k_('α'), t('·'), sub(t('h'), t('1%')), t(' ;     '), sub(t('h'), t('1%')), t(' = '),
                     sub(t('K'), t('1%')), t('·'), gach(t('h'))], '2-10')
    doan(d, 'Trong đó: V – vận tốc gió tính toán ở độ cao 10 m trên mặt nước, V = V_{w} = k_{fl}·k_{l}·V_{l} (TCVN 8421:2010, '
            'Phụ lục A, A.3.3), với V_{l} – vận tốc gió quan trắc ở trạm Phan Thiết (Bảng 1.4), k_{l} = 1,10 ÷ 1,09 – hệ số '
            'quy đổi về mặt nước cho địa hình dạng A (vùng ven biển trống trải, Bảng A3), k_{fl} – hệ số quy đổi số liệu '
            'phong tiêu, lấy bằng 1 (thiên an toàn); D – đà gió (m); H – chiều sâu nước trước đập (m); α – góc giữa hướng '
            'gió và pháp tuyến mái thượng lưu; h̄, τ̄ – chiều cao và chu kỳ sóng trung bình, sóng nước sâu khi H > 0,5λ̄ với '
            'bước sóng λ̄ = g·τ̄^{2}/(2π), lấy giá trị nhỏ hơn giữa điều kiện đà gió và điều kiện thời gian gió thổi t = 6 '
            'giờ; K_{1%} – hệ số theo gD/V^{2} (Hình 2 của tiêu chuẩn); k_{r}, k_{p} – hệ số nhám và thấm của lớp gia cố '
            'mái theo độ nhám tương đối Δ/h_{1%} (mái gia cố đá xây, Δ = 0,02 m); k_{sp} – hệ số theo vận tốc gió và hệ '
            'số mái (m = 3); k_{run} – hệ số theo độ thoải sóng λ̄/h_{1%} (Hình 10); k_{α} – hệ số giảm sóng leo theo góc α '
            '(Bảng 9: 0° – 1,00; 20° – 0,96; 40° – 0,87; 60° – 0,76).', thut=1.0)

    tieu_de(d, '2.3.2 Kết quả tính toán')
    doan(d, 'Mực nước tính toán: MNDBT = +%s m; MNLTK = +%s m và MNLKT = +%s m – kết quả điều tiết lũ khi các cửa van '
            'vận hành bình thường (mục 2.2.3). Tuyến đập có phương vị N96,2°E nên pháp tuyến mái thượng lưu hướng N6,2°E '
            '(lòng hồ ở phía Bắc). Sóng được tính cho các hướng gió thổi vào mái thượng lưu với góc tới không quá 60°: '
            'Bắc (%s°), Đông Bắc (%s°) và Tây Bắc (%s°); các hướng Đông, Tây gần song song với tuyến đập (góc tới 84 ÷ 96°), '
            'các hướng Nam thổi từ phía hạ lưu nên không gây sóng leo lên mái thượng lưu. Đà gió lấy D = %s m khi hồ ở '
            'MNDBT và %s m khi hồ ở MNLTK, dùng chung cho mọi hướng (thiên an toàn); cao trình đáy hồ trước đập +59,0 m. '
            'Kết quả tính chi tiết với gió hướng Bắc cho trong Bảng 2.4, tổng hợp theo hướng gió trong Bảng 2.5.'
            % (so(MNDBT), so(MNLTK), so(MNLKT), so(KQ['N']['a'], 1), so(KQ['NE']['a'], 1), so(KQ['NW']['a'], 1),
               '{:,.0f}'.format(D1).replace(',', '.'), '{:,.0f}'.format(D2).replace(',', '.')), thut=1.0)

    n1, n2 = KQ['N']['r1'], KQ['N']['r2']
    du_lieu = [
        ['Mực nước tính toán (m)', '+' + so(MNDBT), '+' + so(MNLTK), '+' + so(MNLKT)],
        ['Vận tốc gió V_{l} / V_{w} (m/s)', '%s / %s (P = 4%%)' % (so(n1['V_tram'], 1), so(n1['V'])),
         '%s / %s (P = 50%%)' % (so(n2['V_tram'], 1), so(n2['V'])), '–'],
        ['Góc tới α (°) ; k_{α}', '%s ; %s' % (so(KQ['N']['a'], 1), so(n1['ka'])), '%s ; %s' % (so(KQ['N']['a'], 1),
                                                                                         so(n2['ka'])), '–'],
        ['Đà gió D (m)', '{:,.0f}'.format(D1).replace(',', '.'), '{:,.0f}'.format(D2).replace(',', '.'), '–'],
        ['Chiều sâu nước trước đập H (m)', so(n1['H']), so(n2['H']), '–'],
        ['gD/V_{w}^{2}', so(n1['gD'], 1), so(n2['gD'], 1), '–'],
        ['Chiều cao sóng trung bình h̄ (m)', so(n1['h'], 3), so(n2['h'], 3), '–'],
        ['Chu kỳ τ̄ (s) / bước sóng λ̄ (m)', '%s / %s' % (so(n1['tau']), so(n1['lam'])),
         '%s / %s' % (so(n2['tau']), so(n2['lam'])), '–'],
        ['K_{1%} ; h_{1%} (m)', '%s ; %s' % (so(n1['K1']), so(n1['h1'], 3)), '%s ; %s' % (so(n2['K1']), so(n2['h1'], 3)), '–'],
        ['k_{r} ; k_{p} ; k_{sp} ; k_{run}', ' ; '.join(so(n1[k]) for k in ('kr', 'kp', 'ksp', 'krun')),
         ' ; '.join(so(n2[k]) for k in ('kr', 'kp', 'ksp', 'krun')), '–'],
        ['Sóng leo h_{sl1%} (m)', so(n1['hsl'], 3), so(n2['hsl'], 3), '–'],
        ['Độ dềnh do gió Δh (m)', so(n1['dh'], 3), so(n2['dh'], 3), '–'],
        ['Độ vượt cao an toàn (m)', so(A1), so(A2), so(A3)],
        ['Cao trình đỉnh yêu cầu (m)', '+' + so(KQ['N']['Z1']), '+' + so(KQ['N']['Z2']), '+' + so(KQ['N']['Z3'])],
    ]
    bang(d, 'Bảng 2.4 Kết quả tính cao trình đỉnh đập theo ba trường hợp – gió hướng Bắc',
         ['Đại lượng', 'Z_{1}: MNDBT', 'Z_{2}: MNLTK', 'Z_{3}: MNLKT'], du_lieu, cot=[3371, 2100, 2100, 1500])

    hang = []
    for ma, ten, _ in HUONG:
        x = KQ[ma]
        hang.append([ten, so(x['a'], 1), so(x['r1']['ka']), '%s / %s' % (so(x['r1']['V'], 1), so(x['r2']['V'], 1)),
                     '%s / %s' % (so(x['r1']['hsl']), so(x['r2']['hsl'])), so(x['Z1']), so(x['Z2']), so(x['Z3']),
                     so(x['Zd'])])
    bang(d, 'Bảng 2.5 Cao trình đỉnh đập yêu cầu theo hướng gió (m)',
         ['Hướng gió', 'Góc tới α (°)', 'k_{α}', 'V_{w} P = 4% / 50% (m/s)', 'h_{sl} / h′_{sl} (m)', 'Z_{1}', 'Z_{2}',
          'Z_{3}', 'Z_{đ} = max'], hang, cot=[1271, 850, 700, 1500, 1350, 850, 850, 850, 850])
    ne, nb = KQ[MA_MAX], KQ['N']
    doan(d, 'Trị số lớn nhất là Z_{2} = +%s m ứng với gió hướng %s: dù góc tới lớn hơn (k_{α} = %s), vận tốc gió bình quân '
            'lớn nhất hướng %s lớn hơn hướng Bắc (V_{50%%} = %s so với %s m/s) nên sóng leo lớn hơn (h′_{sl} = %s so với '
            '%s m). Như vậy cao trình đỉnh đập yêu cầu Z_{đ} = max(Z_{1}; Z_{2}; Z_{3}) = +%s m.'
            % (so(ZD), ne['ten'], so(ne['r2']['ka']), ne['ten'], so(ne['r2']['V_tram'], 1), so(nb['r2']['V_tram'], 1),
               so(ne['r2']['hsl']), so(nb['r2']['hsl']), so(ZD)), thut=1.0)

    thieu_ht = Z_LKT03 - Z_DINH_HT
    dk = [
        ['Đỉnh tường chắn sóng ≥ Z_{đ} = max(Z_{1}; Z_{2}; Z_{3})', '≥ +' + so(ZD), '+' + so(Z_TUONG),
         'Đạt' if Z_TUONG >= ZD else 'Không đạt', '+' + so(Z_TUONG), 'Đạt' if Z_TUONG >= ZD else 'Không đạt'],
        ['Đỉnh đập (không kể tường) ≥ MNLKT + 0,3 m', '≥ +' + so(Z_LKT03), '+' + so(Z_DINH_HT),
         ('Không đạt (thiếu %s m)' % so(thieu_ht)) if thieu_ht > 0 else 'Đạt', '+' + so(Z_DINH_SC),
         ('Đạt (dư %s m)' % so(Z_DINH_SC - Z_LKT03)) if Z_DINH_SC >= Z_LKT03 else 'Không đạt'],
        ['Tường cao hơn đỉnh đập ≤ 1,2 m', '≤ 1,20', so(Z_TUONG - Z_DINH_HT), 'Đạt' if Z_TUONG - Z_DINH_HT <= 1.2 else
         'Không đạt', so(Z_TUONG - Z_DINH_SC), 'Đạt' if Z_TUONG - Z_DINH_SC <= 1.2 else 'Không đạt'],
    ]
    bang(d, 'Bảng 2.6 Đánh giá cao trình đỉnh đập theo TCVN 8216:2018 điều 10.1.5 (m)',
         ['Điều kiện', 'Yêu cầu', 'Hiện trạng', 'Đánh giá', 'Sau sửa chữa', 'Đánh giá'], dk,
         cot=[3171, 1150, 1100, 1350, 1150, 1150])
    doan(d, 'Đập Cà Giây có tường chắn sóng thẳng đứng phía thượng lưu (đỉnh +%s m) liên kết với bộ phận chống thấm nên '
            'theo TCVN 8216:2018 điều 10.1.5 độ vượt cao được tính đến đỉnh tường: +%s m ≥ Z_{đ} = +%s m – tường chắn sóng '
            'hiện trạng đủ chắn sóng (dự trữ %s m); đồng thời đỉnh đập (không kể tường) phải cao hơn MNLKT tối thiểu 0,3 m, '
            'tức không thấp hơn +%s m, và tường không cao hơn đỉnh đập quá 1,2 m (Bảng 2.6).'
            % (so(Z_TUONG), so(Z_TUONG), so(ZD), so(Z_TUONG - ZD), so(Z_LKT03)), thut=1.0)
    doan(d, 'Cao trình đỉnh đập hiện trạng +%s m (theo hồ sơ thiết kế; khảo sát đo được +%s ÷ +%s m) thấp hơn yêu cầu %s m '
            '(%s ÷ %s m theo số liệu khảo sát): đập hiện trạng chưa bảo đảm an toàn phòng lũ về cao trình đỉnh, cần nâng '
            'đỉnh đập lên tối thiểu +%s m. Với cao trình đỉnh đập sau sửa chữa +%s m: cao hơn MNLKT %s m (≥ 0,3 m), tường '
            'chắn sóng cao hơn đỉnh đập %s m (≤ 1,2 m) – đạt yêu cầu. Vì Z_{đ} = +%s m cao hơn đỉnh đập +%s m nên tường chắn '
            'sóng là bộ phận bắt buộc, phải được giữ lại và duy trì khi sửa chữa.'
            % (so(Z_DINH_HT), so(Z_DINH_KS[0]), so(Z_DINH_KS[1]), so(thieu_ht), so(Z_LKT03 - Z_DINH_KS[1]),
               so(Z_LKT03 - Z_DINH_KS[0]), so(Z_LKT03), so(Z_DINH_SC), so(Z_DINH_SC - MNLKT), so(Z_TUONG - Z_DINH_SC),
               so(ZD), so(Z_DINH_SC)), thut=1.0)
    d.save(RA)
    print('Đã ghi', os.path.abspath(RA))
    for ma in KQ:
        x = KQ[ma]
        print(ma, 'α %.1f Z1 %.3f Z2 %.3f Z3 %.3f Zd %.3f' % (x['a'], x['Z1'], x['Z2'], x['Z3'], x['Zd']))


if __name__ == '__main__':
    main()
