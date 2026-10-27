# -*- coding: utf-8 -*-
"""
XUẤT FILE EXCEL: XÁC ĐỊNH CAO TRÌNH ĐỈNH ĐẬP CÀ GIÂY (mục 2.3 thuyết minh của SV) – TCVN 8216:2018, TCVN 8421:2010
Cùng phương pháp với cao_trinh_dinh.py (đã đối chiếu toàn văn hai tiêu chuẩn đợt 6), mực nước theo yêu cầu SV
27/10/2026: MNDBT +74,70; MNLTK +77,43, MNLKT +78,39 (điều tiết lũ QTVH 72 giờ, vận hành bình thường – file
DieuTietLu_Potapov_f1f2_HoCaGiay.xlsx, sheet KetQua).
  Z1 = MNDBT + Δh + hsl + a (gió P = 4%);  Z2 = MNLTK + Δh′ + h′sl + a′ (gió P = 50%);  Z3 = MNLKT + a″
  Δh = kw·V²·D/(g·H)·cos α, kw tra Bảng A2 (V ≤ 20 m/s: 2,1·10⁻⁶);  sóng nước sâu: TCVN 8421:2010 chỉ cho đồ thị
  (Hình 1, đường bao trên) – Excel dùng công thức xấp xỉ đường bao theo đà gió (Krưlốp); điều kiện thời gian gió
  thổi t = 6 giờ không khống chế (gt/V ≈ 10⁴), có ô để nhập giá trị đọc đồ thị nếu cần (sửa 27/10 – bỏ công thức quy
  đổi thời gian → đà gió tự lập, không có nguồn);
  hsl1% = kr·kp·ksp·krun·kα·h1%, h1% = K1%·h̄;  V = Vw = kfl·kl·Vl (Phụ lục A, A.3.3)
Tính cho các hướng gió thổi vào mái thượng lưu với góc tới ≤ 60°: Bắc, Đông Bắc, Tây Bắc (đà gió lớn nhất cho mọi
hướng). Mọi ô kết quả là công thức; ô nền vàng chữ xanh là số liệu nhập; bảng tra đọc từ bảng, đồ thị của tiêu chuẩn.
  python3 xuat_excel_cao_trinh_dinh.py [đường dẫn .xlsx ra]
Mặc định: ../TAI_VE/CaoTrinhDinhDap_HoCaGiay.xlsx
"""
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import so_lieu as S

DIR = os.path.dirname(os.path.abspath(__file__))
RA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DIR, '..', 'TAI_VE', 'CaoTrinhDinhDap_HoCaGiay.xlsx')
MN = dict(MNDBT=74.70, MNLTK=77.43, MNLKT=78.39)

FONT = 'Times New Roman'
f_thuong = Font(name=FONT, size=11)
f_dam = Font(name=FONT, size=11, bold=True)
f_tieu_de = Font(name=FONT, size=14, bold=True, color='12324A')
f_muc = Font(name=FONT, size=12, bold=True, color='A65A10')
f_nhap = Font(name=FONT, size=11, color='0000FF')
f_nhap_dam = Font(name=FONT, size=11, color='0000FF', bold=True)
f_trang = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_ghi_chu = Font(name=FONT, size=10, italic=True, color='52514E')
nen_nhap = PatternFill('solid', fgColor='FFF2CC')
nen_tieu_de = PatternFill('solid', fgColor='12324A')
nen_kq = PatternFill('solid', fgColor='FBF3E8')
nen_dat = PatternFill('solid', fgColor='E3F4EA')
nen_khong = PatternFill('solid', fgColor='F8D7D7')
vien = Border(*(Side(style='thin', color='B8BEC6'),) * 4)
giua = Alignment(horizontal='center', vertical='center', wrap_text=True)
trai = Alignment(horizontal='left', vertical='center', wrap_text=True)
cg = Alignment(horizontal='center', vertical='center')

wb = Workbook()


def ghi(ws, o, gt, font=f_thuong, fmt=None, fill=None, can=None, border=True):
    c = ws[o]
    c.value = gt
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if can:
        c.alignment = can
    if border:
        c.border = vien
    return c


def ten(name, ref):
    # Excel không phân biệt chữ hoa, chữ thường trong tên (a_1 và A_1 là một tên)
    assert all(n.upper() != name.upper() for n in wb.defined_names), name
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def tieu_de_cot(ws, hang, cot0, nhan, cao=None):
    for k, t in enumerate(nhan):
        ghi(ws, '%s%d' % (get_column_letter(cot0 + k), hang), t, f_trang, fill=nen_tieu_de, can=giua)
    if cao:
        ws.row_dimensions[hang].height = cao


def in_an(ws, ngang=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if ngang else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def noi_suy(x, xs, ys):
    """Nội suy tuyến tính y(x) trong bảng tra (xs tăng dần), ngoài khoảng lấy giá trị đầu, cuối bảng."""
    k = 'MATCH({x},{xs},1)'.format(x=x, xs=xs)
    return ('IF({x}<=INDEX({xs},1),INDEX({ys},1),IF({x}>=INDEX({xs},COUNT({xs})),INDEX({ys},COUNT({xs})),'
            'INDEX({ys},{k})+({x}-INDEX({xs},{k}))*(INDEX({ys},{k}+1)-INDEX({ys},{k}))/(INDEX({xs},{k}+1)-INDEX({xs},{k}))))'
            .format(x=x, xs=xs, ys=ys, k=k))


# ================================================================== HƯỚNG DẪN
hd = wb.active
hd.title = 'HuongDan'
hd.column_dimensions['A'].width = 3
hd.column_dimensions['B'].width = 125
hd.sheet_view.showGridLines = False
ND = [
    ('T', 'XÁC ĐỊNH CAO TRÌNH ĐỈNH ĐẬP HỒ CÀ GIÂY – TCVN 8216:2018 VÀ TCVN 8421:2010'),
    ('g', 'Mọi ô kết quả là công thức. Sửa mực nước, gió, đà gió, cao trình công trình ở sheet SoLieu (ô nền vàng chữ xanh).'),
    ('M', '1. Ba trường hợp (TCVN 8216:2018)'),
    ('c', 'Z₁ = MNDBT + Δh + hsl + a          Z₂ = MNLTK + Δh′ + h′sl + a′          Z₃ = MNLKT + a″'),
    ('-', 'Gió tính sóng (Bảng 4, công trình cấp II): ở MNDBT – vận tốc gió lớn nhất tần suất P = 4%; ở MNLTK – vận tốc gió '
          'bình quân lớn nhất P = 50%; ở MNLKT không xét sóng do gió. Độ vượt cao an toàn (Bảng 3, cấp II): a = 0,7 m; '
          'a′ = 0,5 m; a″ = 0,2 m. Cao trình đỉnh đập yêu cầu Zđ = max(Z₁; Z₂; Z₃).'),
    ('-', 'Điều 10.1.5 – đập có tường chắn sóng thẳng đứng liên kết với bộ phận chống thấm: độ vượt cao tính đến đỉnh tường '
          '(đỉnh tường ≥ Zđ); đỉnh đập không kể tường ≥ MNLKT + 0,3 m; tường cao hơn đỉnh đập không quá 1,2 m.'),
    ('M', '2. Độ dềnh do gió và sóng leo (TCVN 8421:2010)'),
    ('c', 'Δh = kw·V²·D/(g·H)·cos α ,   kw tra Bảng A2: V = 20; 30; 40; 50 m/s → kw = 2,1; 3,0; 3,9; 4,8·10⁻⁶'),
    ('c', 'gh̄/V² = 0,16·[1 − (1 + 0,006·√(gD/V²))⁻²] ;   gτ̄/V = 19,1·(gh̄/V²)^0,625 ;   λ̄ = gτ̄²/(2π)'),
    ('c', 'hsl1% = kr·kp·ksp·krun·kα·h1% ;   h1% = K1%·h̄'),
    ('-', 'V = Vw = kfl·kl·Vl – vận tốc gió tính toán ở độ cao 10 m trên mặt nước (Phụ lục A, A.3.3): Vl – gió trạm Phan Thiết '
          '(Bảng 1.4 thuyết minh), kl – Bảng A3 (địa hình dạng A, vùng ven biển trống trải), kfl = 1 (thiên an toàn).'),
    ('-', 'Sóng nước sâu (H > 0,5λ̄): TCVN 8421:2010 cho gh̄/V² và gτ̄/V bằng đồ thị (Hình 1, đường bao trên), tra theo gD/V² '
          '(điều kiện đà gió) và gt/V (điều kiện thời gian gió thổi), lấy giá trị nhỏ hơn – tiêu chuẩn không cho công thức. '
          'Để Excel tự tính, đường bao theo đà gió được thay bằng công thức xấp xỉ của Krưlốp (hai công thức dòng trên, '
          'λ̄ = gτ̄²/(2π) là công thức của tiêu chuẩn) – nên đọc thử đồ thị 1 ÷ 2 điểm để đối chiếu. Với t = 6 giờ, gt/V ≈ 10⁴: '
          'sóng gần phát triển hoàn toàn, đồ thị cho gh̄/V² lớn hơn nhiều so với điều kiện đà gió nên điều kiện thời gian '
          'không khống chế; ô “theo điều kiện thời gian” ở TinhSong để trống, nhập giá trị đọc đồ thị nếu muốn kiểm tra.'),
    ('-', 'K1% – đồ thị Hình 2 theo gD/V²; kr, kp – bảng hệ số nhám, thấm theo Δ/h1% (mái gia cố đá xây, Δ = 0,02 m); ksp – '
          'bảng theo V và hệ số mái (m = 3 ÷ 5); krun – đồ thị Hình 10 theo λ̄/h1% (đường m = 3); kα – Bảng 9 theo góc tới α '
          'giữa hướng gió và pháp tuyến mái. Các bảng tra ở cuối sheet SoLieu (đồ thị đã đọc thành bảng, nội suy tuyến tính).'),
    ('M', '3. Hướng gió'),
    ('-', 'Tuyến đập phương vị N96,2°E → pháp tuyến mái thượng lưu N6,2°E (lòng hồ phía Bắc). Tính các hướng có góc tới '
          '≤ 60°: Bắc (6,2°), Đông Bắc (38,8°), Tây Bắc (51,2°); Đông, Tây gần song song tuyến đập, các hướng Nam thổi từ '
          'hạ lưu – không gây sóng leo lên mái thượng lưu. Đà gió lớn nhất dùng cho mọi hướng (thiên an toàn).'),
    ('M', '4. Trình tự trong file'),
    ('-', 'SoLieu: mực nước, cấp công trình, gió, đà gió, cao trình đỉnh đập và tường chắn sóng, các bảng tra.'),
    ('-', 'TinhSong: tính từng bước cho 6 trường hợp (3 hướng gió × MNDBT, MNLTK) – các cột cùng công thức.'),
    ('-', 'KetQua: Z₁, Z₂, Z₃ theo hướng gió, Zđ = max, đánh giá đỉnh đập hiện trạng và sau sửa chữa theo điều 10.1.5; '
          'bảng in theo mẫu Bảng 2.4 (gió Bắc).'),
]
for h, (loai, txt) in enumerate(ND, 1):
    o = hd['B%d' % h]
    o.value = ('•  ' + txt) if loai == '-' else txt
    o.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
    if loai == 'T':
        o.font, hd.row_dimensions[h].height = f_tieu_de, 24
    elif loai == 'M':
        o.font, hd.row_dimensions[h].height = f_muc, 22
    elif loai == 'c':
        o.font = Font(name=FONT, size=12, bold=True, color='12324A')
        o.alignment = Alignment(horizontal='center', vertical='center')
        hd.row_dimensions[h].height = 22
    elif loai == 'g':
        o.font, hd.row_dimensions[h].height = f_ghi_chu, 18
    else:
        o.font = f_thuong
        hd.row_dimensions[h].height = 15.5 * max(1, math.ceil(len(txt) / 145))
in_an(hd, ngang=True)

# ================================================================== SỐ LIỆU
sl = wb.create_sheet('SoLieu')
for c, w in zip('ABCDEFGHI', (46, 11, 9, 11, 11, 11, 11, 11, 11)):
    sl.column_dimensions[c].width = w
ghi(sl, 'A1', 'SỐ LIỆU ĐẦU VÀO – CAO TRÌNH ĐỈNH ĐẬP', f_tieu_de, border=False)
ghi(sl, 'A2', 'Ô nền vàng chữ xanh: được sửa. Chữ đen: công thức.', f_ghi_chu, border=False)


def khoi(h0, tieu_de, dong):
    ghi(sl, 'A%d' % h0, tieu_de, f_muc, border=False)
    tieu_de_cot(sl, h0 + 1, 1, ['Thông số', 'Giá trị', 'Đơn vị'])
    ghi(sl, 'D%d' % (h0 + 1), 'Ghi chú, nguồn', f_trang, fill=nen_tieu_de, can=giua)
    sl.merge_cells('D%d:I%d' % (h0 + 1, h0 + 1))
    for k, (nhan, gt, dv, ten_o, gc, fmt) in enumerate(dong):
        r = h0 + 2 + k
        nhap = not (isinstance(gt, str) and gt.startswith('='))
        ghi(sl, 'A%d' % r, nhan, f_dam)
        ghi(sl, 'B%d' % r, gt, f_nhap_dam if nhap else f_dam, fmt, nen_nhap if nhap else nen_kq, cg)
        ghi(sl, 'C%d' % r, dv, can=cg)
        ghi(sl, 'D%d' % r, gc, f_ghi_chu)
        sl.merge_cells('D%d:I%d' % (r, r))
        if ten_o:
            ten(ten_o, 'SoLieu!$B$%d' % r)
    return h0 + 2 + len(dong)


h = khoi(4, 'A. MỰC NƯỚC HỒ', [
    ('Mực nước dâng bình thường MNDBT', MN['MNDBT'], 'm', 'MN_DBT', '', '0.00'),
    ('Mực nước lũ thiết kế MNLTK (P = 1%)', MN['MNLTK'], 'm', 'MN_LTK',
     'Điều tiết lũ QTVH 72 giờ, vận hành bình thường (DieuTietLu_Potapov – KetQua)', '0.00'),
    ('Mực nước lũ kiểm tra MNLKT (P = 0,2%)', MN['MNLKT'], 'm', 'MN_LKT', 'Như trên', '0.00'),
    ('Cao trình đáy hồ trước đập', 59.0, 'm', 'Z_day', 'Tính chiều sâu nước H = Z − Z đáy (H = 15,70 m ở MNDBT)', '0.00'),
    ('Gia tốc trọng trường g', S.G, 'm/s²', 'g_', '', '0.00'),
])
h = khoi(h + 1, 'B. ĐỘ VƯỢT CAO AN TOÀN – TCVN 8216:2018 Bảng 3, đập cấp II', [
    ('a – trường hợp MNDBT', 0.70, 'm', 'a_1', '', '0.00'),
    ('a′ – trường hợp MNLTK', 0.50, 'm', 'a_2', '', '0.00'),
    ('a″ – trường hợp MNLKT', 0.20, 'm', 'a_3', '', '0.00'),
])
h = khoi(h + 1, 'C. GIÓ, SÓNG VÀ MÁI ĐẬP', [
    ('Thời gian gió thổi t', 6, 'giờ', 't_gio', 'Điều kiện thời gian khi tính sóng', '0'),
    ('Hệ số quy đổi số liệu phong tiêu kfl', 1.0, '', 'k_fl', 'TCVN 8421:2010 A.3.3 – lấy bằng 1 (thiên an toàn)', '0.00'),
    ('Đà gió D khi hồ ở MNDBT', 2970, 'm', 'D_1', 'Như Bảng 2.4 (bản trước); dùng cho mọi hướng – thiên an toàn', '0'),
    ('Đà gió D khi hồ ở MNLTK', 3025, 'm', 'D_2', 'Như trên (ứng với MNLTK +77,43 của hồ sơ BCNCKT)', '0'),
    ('Hệ số mái thượng lưu m', 3.0, '', 'm_TL', 'Bảng ksp dưới đây dùng hàng m = 3 ÷ 5', '0.0'),
    ('Độ nhám lớp gia cố mái Δ', 0.02, 'm', 'Delta_n', 'Mái gia cố đá xây', '0.000'),
    ('Phương vị pháp tuyến mái thượng lưu', 6.2, '°', 'PT_mai', 'Tuyến đập N96,2°E, lòng hồ phía Bắc', '0.0'),
])
h = khoi(h + 1, 'D. CAO TRÌNH CÔNG TRÌNH (đánh giá theo TCVN 8216:2018 điều 10.1.5)', [
    ('Đỉnh đập hiện trạng (hồ sơ thiết kế)', 78.50, 'm', 'Z_dinh_HT', 'Khảo sát đo +78,32 ÷ +78,44 – chỉ để đối chiếu', '0.00'),
    ('Đỉnh tường chắn sóng', 79.40, 'm', 'Z_tuong', 'Tường chắn sóng thẳng đứng phía thượng lưu – giữ nguyên', '0.00'),
    ('Đỉnh đập sau sửa chữa', S.Z_DINH_DAP_SC, 'm', 'Z_dinh_SC', 'Phương án chọn của đồ án', '0.00'),
    ('Đỉnh đập không kể tường cao hơn MNLKT tối thiểu', 0.30, 'm', 'dZ_LKT', 'Điều 10.1.5', '0.00'),
    ('Tường cao hơn đỉnh đập tối đa', 1.20, 'm', 'dZ_tuong', 'Điều 10.1.5', '0.00'),
])
# E. gió theo hướng
HG = h + 1
ghi(sl, 'A%d' % HG, 'E. VẬN TỐC GIÓ LỚN NHẤT THEO HƯỚNG – TRẠM PHAN THIẾT (Bảng 1.4 thuyết minh), m/s', f_muc, border=False)
ghi(sl, 'H%d' % HG, 'Tần suất gió theo TCVN 8216:2018 Bảng 4 (cấp II): MNDBT – P = 4%; MNLTK – P = 50% (gió bình quân '
                    'lớn nhất); MNLKT – không xét sóng.', f_ghi_chu, border=False)
tieu_de_cot(sl, HG + 1, 1, ['Hướng gió', 'Mã', 'Phương vị (°)', 'V P = 4%', 'V P = 50%', 'Góc tới α (°)',
                            'Thổi vào mái TL (α ≤ 60°)?'], cao=32)
TEN_HUONG = dict(N='Bắc', NE='Đông Bắc', E='Đông', SE='Đông Nam', S='Nam', SW='Tây Nam', W='Tây', NW='Tây Bắc')
AZ = dict(N=0, NE=45, E=90, SE=135, S=180, SW=225, W=270, NW=315)
for k, ma in enumerate(['N', 'NE', 'E', 'SE', 'S', 'SW', 'W', 'NW']):
    r = HG + 2 + k
    j = S.GIO_HUONG.index(ma)
    ghi(sl, 'A%d' % r, TEN_HUONG[ma], f_dam)
    ghi(sl, 'B%d' % r, ma, can=cg)
    ghi(sl, 'C%d' % r, AZ[ma], fmt='0', can=cg)
    ghi(sl, 'D%d' % r, S.GIO[4][j], f_nhap, '0.0', nen_nhap, cg)
    ghi(sl, 'E%d' % r, S.GIO[50][j], f_nhap, '0.0', nen_nhap, cg)
    ghi(sl, 'F%d' % r, '=ABS(MOD(C{r}-PT_mai+180,360)-180)'.format(r=r), fmt='0.0', fill=nen_kq, can=cg)
    ghi(sl, 'G%d' % r, '=IF(F{r}<=60,"Có – tính","Không")'.format(r=r), can=cg)
RG0, RG1 = HG + 2, HG + 9
ten('Gio_ma', 'SoLieu!$B$%d:$B$%d' % (RG0, RG1))
ten('Gio_az', 'SoLieu!$C$%d:$C$%d' % (RG0, RG1))
ten('Gio_V4', 'SoLieu!$D$%d:$D$%d' % (RG0, RG1))
ten('Gio_V50', 'SoLieu!$E$%d:$E$%d' % (RG0, RG1))

# F. bảng tra
HT = RG1 + 3
ghi(sl, 'A%d' % HT, 'F. BẢNG TRA (TCVN 8421:2010) – đồ thị đã đọc thành bảng, nội suy tuyến tính, ngoài bảng lấy giá trị biên',
    f_muc, border=False)
BANG_TRA = [
    ('Hệ số kl quy đổi gió về mặt nước, địa hình dạng A (Bảng A3)', 'Vl (m/s)', 'kl', 'KL_V', 'KL_k',
     [10, 15, 20, 25, 30, 35, 40], [1.10, 1.10, 1.09, 1.09, 1.09, 1.09, 1.08], '0', '0.00'),
    ('Hệ số kα theo góc tới α (Bảng 9)', 'α (°)', 'kα', 'KA_goc', 'KA_k',
     [0, 10, 20, 30, 40, 50, 60], [1.00, 0.98, 0.96, 0.92, 0.87, 0.82, 0.76], '0', '0.00'),
    ('Hệ số nhám kr theo Δ/h1% (bảng hệ số nhám, thấm)', 'Δ/h1%', 'kr', 'KR_x', 'KR_kr',
     [0.002, 0.0075, 0.02, 0.05, 0.1, 0.2], [1.00, 0.95, 0.90, 0.80, 0.75, 0.70], '0.0000', '0.00'),
    ('Hệ số thấm kp theo Δ/h1% (cùng bảng)', 'Δ/h1%', 'kp', 'KP_x', 'KP_kp',
     [0.002, 0.0075, 0.02, 0.05, 0.1, 0.2], [0.90, 0.85, 0.80, 0.70, 0.60, 0.50], '0.0000', '0.00'),
    ('Hệ số ksp theo V, hàng mái m = 3 ÷ 5 (m < 3: 1,1 / 1,1 / 1,4; m > 5: 0,6 / 1,2 / 1,6)', 'V (m/s)', 'ksp', 'KSP_V',
     'KSP_k', [5, 10, 20], [0.8, 1.1, 1.5], '0', '0.00'),
    ('Hệ số K1% theo gD/V² (đồ thị Hình 2)', 'gD/V²', 'K1%', 'K1_x', 'K1_k',
     [20, 100, 153, 514, 1000, 5000], [2.00, 2.08, 2.10, 2.13, 2.15, 2.18], '0', '0.00'),
    ('Hệ số krun theo λ̄/h1%, mái m = 3 (đồ thị Hình 10)', 'λ̄/h1%', 'krun', 'KRUN_x', 'KRUN_k',
     [5, 8, 10, 12, 15, 20, 25], [1.20, 1.40, 1.48, 1.57, 1.70, 1.84, 1.95], '0', '0.00'),
    ('Hệ số kw tính độ dềnh do gió (Bảng A2; V ≤ 20 m/s lấy 2,1·10⁻⁶ – bảng không có giá trị nhỏ hơn)', 'V (m/s)', 'kw',
     'KW_V', 'KW_k', [20, 30, 40, 50], [2.1e-6, 3.0e-6, 3.9e-6, 4.8e-6], '0', '0.0E+00'),
]
r = HT + 1
for tieu, nx, ny, tx, ty, xs, ys, fx, fy in BANG_TRA:
    ghi(sl, 'A%d' % r, tieu, f_dam, border=False)
    ghi(sl, 'A%d' % (r + 1), nx, can=Alignment(horizontal='right'))
    ghi(sl, 'A%d' % (r + 2), ny, can=Alignment(horizontal='right'))
    for k, (x, y) in enumerate(zip(xs, ys)):
        c = get_column_letter(2 + k)
        ghi(sl, '%s%d' % (c, r + 1), x, f_nhap, fx, nen_nhap, cg)
        ghi(sl, '%s%d' % (c, r + 2), y, f_nhap, fy, nen_nhap, cg)
    cl = get_column_letter(1 + len(xs))
    ten(tx, 'SoLieu!$B$%d:$%s$%d' % (r + 1, cl, r + 1))
    ten(ty, 'SoLieu!$B$%d:$%s$%d' % (r + 2, cl, r + 2))
    if tx == 'K1_x':                 # K1% nội suy theo lg(gD/V²)
        ghi(sl, 'A%d' % (r + 3), 'lg(gD/V²)', can=Alignment(horizontal='right'))
        for k in range(len(xs)):
            c = get_column_letter(2 + k)
            ghi(sl, '%s%d' % (c, r + 3), '=LOG10(%s%d)' % (c, r + 1), fmt='0.000', fill=nen_kq, can=cg)
        ten('K1_lg', 'SoLieu!$B$%d:$%s$%d' % (r + 3, cl, r + 3))
        r += 1
    r += 4
sl.freeze_panes = 'A4'
in_an(sl)

# ================================================================== TÍNH SÓNG – 6 TRƯỜNG HỢP
ts = wb.create_sheet('TinhSong')
ts.column_dimensions['A'].width = 50
ts.column_dimensions['B'].width = 9
for c in 'CDEFGH':
    ts.column_dimensions[c].width = 12
ts.column_dimensions['I'].width = 58
ghi(ts, 'A1', 'TÍNH SÓNG DO GIÓ VÀ CAO TRÌNH ĐỈNH YÊU CẦU Z₁, Z₂ – TỪNG BƯỚC (TCVN 8421:2010)', f_tieu_de, border=False)
ghi(ts, 'A2', 'Các cột cùng công thức; muốn tính hướng khác: sửa mã hướng (dòng 4) theo cột “Mã” ở SoLieu mục E.', f_ghi_chu,
    border=False)
CASES = [('N', 'Z1'), ('N', 'Z2'), ('NE', 'Z1'), ('NE', 'Z2'), ('NW', 'Z1'), ('NW', 'Z2')]
tieu_de_cot(ts, 3, 1, ['Đại lượng', 'Đơn vị'] + ['%s – %s' % (TEN_HUONG[a], 'MNDBT' if b == 'Z1' else 'MNLTK')
                                                for a, b in CASES] + ['Công thức, nguồn'], cao=32)
BUOC = [
    ('huong', 'Mã hướng gió', '', None, None, ''),
    ('th', 'Trường hợp (Z1 – MNDBT, gió P = 4%; Z2 – MNLTK, gió P = 50%)', '', None, None, ''),
    ('Z', 'Mực nước tính toán Z', 'm', '=IF({c}{th}="Z1",MN_DBT,MN_LTK)', '0.00', 'SoLieu mục A'),
    ('H', 'Chiều sâu nước trước đập H', 'm', '={c}{Z}-Z_day', '0.00', 'H = Z − Z đáy'),
    ('Vl', 'Vận tốc gió trạm Vl', 'm/s',
     '=IF({c}{th}="Z1",INDEX(Gio_V4,MATCH({c}{huong},Gio_ma,0)),INDEX(Gio_V50,MATCH({c}{huong},Gio_ma,0)))', '0.0',
     'Gió P = 4% (Z1), P = 50% (Z2) – SoLieu mục E'),
    ('kl', 'Hệ số quy đổi kl', '', '=' + noi_suy('{c}{Vl}', 'KL_V', 'KL_k'), '0.000', 'Bảng A3, địa hình dạng A'),
    ('Vw', 'Vận tốc gió tính toán V = Vw', 'm/s', '=k_fl*{c}{kl}*{c}{Vl}', '0.00', 'Vw = kfl·kl·Vl (A.3.3)'),
    ('alpha', 'Góc tới α giữa hướng gió và pháp tuyến mái', '°',
     '=ABS(MOD(INDEX(Gio_az,MATCH({c}{huong},Gio_ma,0))-PT_mai+180,360)-180)', '0.0', ''),
    ('D', 'Đà gió D', 'm', '=IF({c}{th}="Z1",D_1,D_2)', '0', 'SoLieu mục C'),
    ('gD', 'gD/V²', '', '=g_*{c}{D}/{c}{Vw}^2', '0.0', ''),
    ('gt', 'gt/V (t = 6 giờ)', '', '=g_*t_gio*3600/{c}{Vw}', '0', ''),
    ('aD', 'gh̄/V² theo điều kiện đà gió (đường bao Hình 1)', '', '=0.16*(1-(1+0.006*SQRT({c}{gD}))^-2)', '0.0000',
     'Công thức xấp xỉ đường bao (Krưlốp): 0,16·[1 − (1 + 0,006·√(gD/V²))⁻²]'),
    ('at', 'gh̄/V² theo điều kiện thời gian gió thổi – tra Hình 1 với gt/V (không bắt buộc)', '', None, '0.0000',
     'Để trống: gt/V ≈ 10⁴ – sóng gần phát triển hoàn toàn, không khống chế'),
    ('a', 'gh̄/V² chọn (giá trị nhỏ hơn)', '', '=IF(ISNUMBER({c}{at}),MIN({c}{aD},{c}{at}),{c}{aD})', '0.0000', ''),
    ('h', 'Chiều cao sóng trung bình h̄', 'm', '={c}{a}*{c}{Vw}^2/g_', '0.000', 'h̄ = (gh̄/V²)·V²/g'),
    ('tau', 'Chu kỳ sóng trung bình τ̄', 's', '=19.1*{c}{a}^0.625*{c}{Vw}/g_', '0.00',
     'Xấp xỉ đường bao (Krưlốp): gτ̄/V = 19,1·(gh̄/V²)^0,625'),
    ('lam', 'Bước sóng trung bình λ̄', 'm', '=g_*{c}{tau}^2/(2*PI())', '0.00', 'λ̄ = g·τ̄²/(2π)'),
    ('sau', 'Kiểm tra sóng nước sâu H > 0,5λ̄', '', '=IF({c}{H}>0.5*{c}{lam},"Đúng","Sai – xem lại")', None, ''),
    ('K1', 'Hệ số K1%', '', '=' + noi_suy('LOG10({c}{gD})', 'K1_lg', 'K1_k'), '0.000', 'Hình 2, theo gD/V²'),
    ('h1', 'Chiều cao sóng h1% = K1%·h̄', 'm', '={c}{K1}*{c}{h}', '0.000', ''),
    ('rel', 'Độ nhám tương đối Δ/h1%', '', '=Delta_n/{c}{h1}', '0.0000', 'Δ = 0,02 m'),
    ('kr', 'Hệ số nhám kr', '', '=' + noi_suy('{c}{rel}', 'KR_x', 'KR_kr'), '0.000', 'Bảng hệ số nhám, thấm'),
    ('kp', 'Hệ số thấm kp', '', '=' + noi_suy('{c}{rel}', 'KP_x', 'KP_kp'), '0.000', 'Như trên'),
    ('ksp', 'Hệ số ksp', '', '=' + noi_suy('{c}{Vw}', 'KSP_V', 'KSP_k'), '0.000', 'Bảng ksp, mái m = 3 ÷ 5'),
    ('lh', 'Độ thoải sóng λ̄/h1%', '', '={c}{lam}/{c}{h1}', '0.00', ''),
    ('krun', 'Hệ số krun', '', '=' + noi_suy('{c}{lh}', 'KRUN_x', 'KRUN_k'), '0.000', 'Hình 10, mái m = 3'),
    ('ka', 'Hệ số kα', '', '=' + noi_suy('{c}{alpha}', 'KA_goc', 'KA_k'), '0.000', 'Bảng 9, theo α'),
    ('hsl', 'Chiều cao sóng leo hsl1%', 'm', '={c}{kr}*{c}{kp}*{c}{ksp}*{c}{krun}*{c}{ka}*{c}{h1}', '0.000',
     'hsl1% = kr·kp·ksp·krun·kα·h1%'),
    ('kw', 'Hệ số kw', '', '=' + noi_suy('{c}{Vw}', 'KW_V', 'KW_k'), '0.00E+00', 'Bảng A2 (V ≤ 20 m/s: 2,1·10⁻⁶)'),
    ('dh', 'Độ dềnh do gió Δh', 'm', '={c}{kw}*{c}{Vw}^2*{c}{D}/(g_*{c}{H})*COS(RADIANS({c}{alpha}))', '0.0000',
     'Δh = kw·V²·D/(g·H)·cos α'),
    ('aa', 'Độ vượt cao an toàn', 'm', '=IF({c}{th}="Z1",a_1,a_2)', '0.00', 'a (Z1), a′ (Z2) – Bảng 3'),
    ('Zyc', 'Cao trình đỉnh đập yêu cầu Z₁ hoặc Z₂', 'm', '={c}{Z}+{c}{dh}+{c}{hsl}+{c}{aa}', '0.00', 'Z = Z + Δh + hsl + a'),
]
R0 = 4
DONG = {k: R0 + i for i, (k, *_) in enumerate(BUOC)}
for i, (k, nhan, dv, ct, fmt, gc) in enumerate(BUOC):
    r = R0 + i
    ghi(ts, 'A%d' % r, nhan, f_dam if k in ('Zyc', 'hsl', 'dh') else f_thuong, can=trai)
    ghi(ts, 'B%d' % r, dv, can=cg)
    ghi(ts, 'I%d' % r, gc, f_ghi_chu, can=trai)
    for j, (huong, th) in enumerate(CASES):
        c = get_column_letter(3 + j)
        if k == 'huong':
            ghi(ts, '%s%d' % (c, r), huong, f_nhap_dam, None, nen_nhap, cg)
        elif k == 'th':
            ghi(ts, '%s%d' % (c, r), th, f_nhap_dam, None, nen_nhap, cg)
        elif k == 'at':
            ghi(ts, '%s%d' % (c, r), None, f_nhap, fmt, nen_nhap, cg)
        else:
            f = ct.format(c=c, **DONG)
            ghi(ts, '%s%d' % (c, r), f, f_dam if k == 'Zyc' else f_thuong, fmt, nen_kq if k in ('hsl', 'dh', 'Zyc') else None, cg)
R_END = R0 + len(BUOC) - 1
ts.freeze_panes = 'C4'
in_an(ts, ngang=True)

# ================================================================== KẾT QUẢ
kq = wb.create_sheet('KetQua')
for c, w in zip('ABCDEFGHIJK', (36, 10, 9, 11, 11, 11, 11, 11, 11, 11, 11)):
    kq.column_dimensions[c].width = w
ghi(kq, 'A1', 'KẾT QUẢ: CAO TRÌNH ĐỈNH ĐẬP YÊU CẦU VÀ ĐÁNH GIÁ (TCVN 8216:2018)', f_tieu_de, border=False)
ghi(kq, 'A2', '="MNDBT +"&TEXT(MN_DBT,"0.00")&"; MNLTK +"&TEXT(MN_LTK,"0.00")&"; MNLKT +"&TEXT(MN_LKT,"0.00")&" m"',
    f_ghi_chu, border=False)
# 1. theo hướng gió
ghi(kq, 'A4', '1. Cao trình đỉnh yêu cầu theo hướng gió (m)', f_muc, border=False)
tieu_de_cot(kq, 5, 1, ['Hướng gió', 'Góc tới α (°)', 'kα', 'Vw P = 4% (m/s)', 'Vw P = 50% (m/s)', 'hsl ở MNDBT (m)',
                       'hsl′ ở MNLTK (m)', 'Z₁', 'Z₂', 'Z₃', 'Zđ = max'], cao=45)
for k, huong in enumerate(['N', 'NE', 'NW']):
    r = 6 + k
    c1, c2 = get_column_letter(3 + 2 * k), get_column_letter(4 + 2 * k)       # cột Z1, Z2 của hướng ở TinhSong
    ghi(kq, 'A%d' % r, TEN_HUONG[huong], f_dam)
    for col, f, fmt in (('B', '=TinhSong!%s%d' % (c1, DONG['alpha']), '0.0'), ('C', '=TinhSong!%s%d' % (c1, DONG['ka']), '0.00'),
                        ('D', '=TinhSong!%s%d' % (c1, DONG['Vw']), '0.00'), ('E', '=TinhSong!%s%d' % (c2, DONG['Vw']), '0.00'),
                        ('F', '=TinhSong!%s%d' % (c1, DONG['hsl']), '0.000'), ('G', '=TinhSong!%s%d' % (c2, DONG['hsl']), '0.000'),
                        ('H', '=TinhSong!%s%d' % (c1, DONG['Zyc']), '0.00'), ('I', '=TinhSong!%s%d' % (c2, DONG['Zyc']), '0.00'),
                        ('J', '=MN_LKT+a_3', '0.00'), ('K', '=MAX(H{r}:J{r})'.format(r=r), '0.00')):
        ghi(kq, '%s%d' % (col, r), f, f_dam if col == 'K' else f_thuong, fmt, nen_kq if col in 'HIJK' else None, cg)
ghi(kq, 'A9', 'Cao trình đỉnh yêu cầu lớn nhất Zđ (m)', f_dam)
ghi(kq, 'B9', '=MAX(K6:K8)', f_dam, '0.00', nen_kq, cg)
ghi(kq, 'C9', '=INDEX(A6:A8,MATCH(B9,K6:K8,0))&" – "&IF(INDEX(I6:I8,MATCH(B9,K6:K8,0))=B9,"Z₂ (MNLTK)",'
              'IF(INDEX(H6:H8,MATCH(B9,K6:K8,0))=B9,"Z₁ (MNDBT)","Z₃ (MNLKT)"))', f_ghi_chu, border=False)
ten('Z_yc', 'KetQua!$B$9')
ghi(kq, 'A10', 'Các hướng Đông, Tây (góc tới 84 ÷ 96°) và các hướng Nam (thổi từ hạ lưu) không gây sóng leo lên mái thượng lưu.',
    f_ghi_chu, border=False)
# 2. đánh giá
ghi(kq, 'A12', '2. Đánh giá theo TCVN 8216:2018 điều 10.1.5 (đập có tường chắn sóng)', f_muc, border=False)
tieu_de_cot(kq, 13, 1, ['Điều kiện', 'Yêu cầu (m)', 'Hiện trạng (m)', 'Chênh (m)', 'Đánh giá', 'Sau sửa chữa (m)', 'Chênh (m)',
                        'Đánh giá'], cao=32)
DK = [
    ('Đỉnh tường chắn sóng ≥ Zđ = max(Z₁; Z₂; Z₃)', '=Z_yc', '=Z_tuong', '=Z_tuong', 'ge'),
    ('Đỉnh đập không kể tường ≥ MNLKT + 0,3 m', '=MN_LKT+dZ_LKT', '=Z_dinh_HT', '=Z_dinh_SC', 'ge'),
    ('Tường cao hơn đỉnh đập ≤ 1,2 m', '=dZ_tuong', '=Z_tuong-Z_dinh_HT', '=Z_tuong-Z_dinh_SC', 'le'),
]
for k, (nhan, yc, ht, sc, kieu) in enumerate(DK):
    r = 14 + k
    ghi(kq, 'A%d' % r, nhan, f_dam, can=trai)
    ghi(kq, 'B%d' % r, yc, fmt='0.00', fill=nen_kq, can=cg)
    for c_gt, c_ch, c_dg, gt in (('C', 'D', 'E', ht), ('F', 'G', 'H', sc)):
        ghi(kq, '%s%d' % (c_gt, r), gt, fmt='0.00', can=cg)
        ghi(kq, '%s%d' % (c_ch, r), ('={g}{r}-B{r}' if kieu == 'ge' else '=B{r}-{g}{r}').format(g=c_gt, r=r),
            fmt='+0.00;-0.00;0.00', can=cg)
        ghi(kq, '%s%d' % (c_dg, r), '=IF({h}{r}>=-0.0001,"Đạt","Không đạt")'.format(h=c_ch, r=r), f_dam, can=cg)
        kq.conditional_formatting.add('%s%d' % (c_dg, r), FormulaRule(formula=['%s%d="Đạt"' % (c_dg, r)], fill=nen_dat))
        kq.conditional_formatting.add('%s%d' % (c_dg, r), FormulaRule(formula=['%s%d="Không đạt"' % (c_dg, r)], fill=nen_khong))
kq.row_dimensions[14].height = kq.row_dimensions[15].height = 30
ghi(kq, 'A18', 'Chênh = giá trị thực − yêu cầu (dòng 1, 2) hoặc yêu cầu − giá trị thực (dòng 3); dương là đạt.', f_ghi_chu,
    border=False)
ghi(kq, 'A19', '="Nếu không có tường chắn sóng, đỉnh đập phải ≥ Zđ = +"&TEXT(Z_yc,"0.00")&" m: hiện trạng "&TEXT(Z_dinh_HT-Z_yc,'
               '"+0.00;-0.00")&" m, sau sửa chữa "&TEXT(Z_dinh_SC-Z_yc,"+0.00;-0.00")&" m → tường chắn sóng phải được giữ lại."',
    f_ghi_chu, border=False)
# 3. bảng in theo mẫu Bảng 2.4 (gió Bắc)
ghi(kq, 'A21', '3. Bảng in theo mẫu Bảng 2.4 – gió hướng Bắc (chép vào thuyết minh)', f_muc, border=False)
tieu_de_cot(kq, 22, 1, ['Đại lượng', 'Z₁: MNDBT', '', 'Z₂: MNLTK', '', 'Z₃: MNLKT'], cao=22)
for a, b in (('B22', 'C22'), ('D22', 'E22')):
    kq.merge_cells('%s:%s' % (a, b))
BANG_IN = [
    ('Mực nước tính toán (m)', 'Z', '0.00', '=MN_LKT'),
    ('Vận tốc gió Vl (m/s)', 'Vl', '0.0', '–'),
    ('Vận tốc gió tính toán Vw (m/s)', 'Vw', '0.00', '–'),
    ('Đà gió D (m)', 'D', '0', '–'),
    ('Chiều sâu nước trước đập H (m)', 'H', '0.00', '–'),
    ('gD/Vw²', 'gD', '0.0', '–'),
    ('Chiều cao sóng trung bình h̄ (m)', 'h', '0.000', '–'),
    ('Chu kỳ τ̄ (s)', 'tau', '0.00', '–'),
    ('Bước sóng λ̄ (m)', 'lam', '0.00', '–'),
    ('K1%', 'K1', '0.00', '–'),
    ('h1% (m)', 'h1', '0.000', '–'),
    ('kr', 'kr', '0.00', '–'), ('kp', 'kp', '0.00', '–'), ('ksp', 'ksp', '0.00', '–'), ('krun', 'krun', '0.00', '–'),
    ('kα', 'ka', '0.00', '–'),
    ('Sóng leo hsl1% (m)', 'hsl', '0.000', '–'),
    ('Độ dềnh Δh (m)', 'dh', '0.000', '–'),
    ('Độ vượt cao an toàn (m)', 'aa', '0.00', '=a_3'),
    ('Cao trình đỉnh yêu cầu (m)', 'Zyc', '0.00', '=MN_LKT+a_3'),
]
for k, (nhan, key, fmt, z3) in enumerate(BANG_IN):
    r = 23 + k
    ghi(kq, 'A%d' % r, nhan, f_dam if key == 'Zyc' else f_thuong, can=trai)
    for c_from, c_to, src in (('B', 'C', 'C'), ('D', 'E', 'D')):
        ghi(kq, '%s%d' % (c_from, r), '=TinhSong!%s%d' % (src, DONG[key]), f_dam if key == 'Zyc' else f_thuong, fmt, can=cg)
        ghi(kq, '%s%d' % (c_to, r), None)
        kq.merge_cells('%s%d:%s%d' % (c_from, r, c_to, r))
    ghi(kq, 'F%d' % r, z3, f_dam if key == 'Zyc' else f_thuong, fmt, can=cg)
kq.freeze_panes = 'A4'
in_an(kq, ngang=True)

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = {'HuongDan': 'A65A10', 'SoLieu': 'EDA100', 'TinhSong': '2A78D6',
                                    'KetQua': '1BAF7A'}[ws.title]
wb.calculation.fullCalcOnLoad = True
wb.save(RA)
print('Đã ghi', os.path.abspath(RA))
