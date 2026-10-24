# -*- coding: utf-8 -*-
"""
XUẤT BẢNG KẾT QUẢ KIỂM TRA ỔN ĐỊNH MÁI ĐẬP SAU SỬA CHỮA RA WORD – MẶT CẮT LÒNG SÔNG VÀ MẶT CẮT SƯỜN ĐỒI
Cùng mẫu Bảng 5.2 (ổn định) của thuyết minh: 12 trường hợp, K hiện trạng (đỉnh +78,50) và K sau sửa chữa (+78,85).
Số liệu: ket_qua/tham_on_dinh.json (chay_tham_on_dinh.py), ket_qua/phan_tich_bo_sung_on_dinh.json
(phan_tich_bo_sung.py – TH11, TH12 thi công), ket_qua/on_dinh_suon_doi.json (on_dinh_suon_doi.py, mục 'chuong5').
Chạy lại các chương trình đó trước nếu thông số đầu vào thay đổi.
  python3 xuat_bang_on_dinh_docx.py
"""
import json
import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

import so_lieu as S
from xuat_bang_tham_docx import bang, doan, so

DIR = os.path.dirname(os.path.abspath(__file__))
RA = os.path.join(DIR, '..', 'TAI_VE', 'BangOnDinh_LongSong_SuonDoi.docx')
HINH_SD = os.path.join(DIR, '..', 'TAI_VE', 'OnDinh_SuonDoi_SauSC.png')
COT = [732, 3831, 563, 676, 901, 901, 676, 789]       # bề rộng cột (dxa) như Bảng 5.2 trong thuyết minh
TIEU_DE = ['TH', 'Trường hợp tính toán', 'Mái', 'Tổ hợp', 'K hiện trạng', 'K sau SC', '[K]', 'Đánh giá']
TO_HOP = {'CB': 'CB', 'DB': 'ĐB', 'TC': 'TC'}
Z_CHAN_LS, Z_CHAN_SD = 54.0, 64.6    # chân mái thượng lưu: lòng sông (tham_fem.MatCat), sườn đồi (tham_suon_doi.MAT)


def dong_bang(ht, sc):
    """ht, sc: {mã: dict(ma, mo_ta, mai, to_hop, K, K_cp)}; hiện trạng không có TH11, TH12."""
    out = []
    for ma, r in sc.items():
        out.append([ma, r['mo_ta'], r['mai'], TO_HOP[r['to_hop']], so(ht[ma]['K'], 3) if ma in ht else '–',
                    so(r['K'], 3), so(r['K_cp']), 'Đạt' if r['K'] >= r['K_cp'] else 'Không đạt'])
    return out


def ten_th(r):
    """Mô tả trường hợp viết thường chữ đầu, giữ nguyên chữ viết tắt (MNDBT, MN…), bỏ phần trong ngoặc."""
    t = r['mo_ta'].split(' (')[0]
    return t[0].lower() + t[1:] if t[1:2].islower() else t


def danh_gia(ht, sc):
    """Các đại lượng cho đoạn nhận xét."""
    khai_thac = [m for m in sc if m in ht]
    hl = min((r for r in sc.values() if r['mai'] == 'HL' and r['to_hop'] != 'TC'), key=lambda r: r['K'])
    tl = min((r for r in sc.values() if r['mai'] == 'TL' and r['to_hop'] != 'TC'), key=lambda r: r['K'])
    thay_doi = [100 * (sc[m]['K'] / ht[m]['K'] - 1) for m in khai_thac]
    du_tru = min(100 * (r['K'] / r['K_cp'] - 1) for r in sc.values())
    return dict(hl=hl, tl=tl, giam_min=-max(thay_doi), giam_max=-min(thay_doi), du_tru=du_tru,
                tc_hl=sc['TH1']['K'] - sc['TH11']['K'], tc_tl=sc['TH7']['K'] - sc['TH12']['K'])


def main():
    T = json.load(open(os.path.join(DIR, 'ket_qua', 'tham_on_dinh.json'), encoding='utf-8'))['on_dinh']
    B = json.load(open(os.path.join(DIR, 'ket_qua', 'phan_tich_bo_sung_on_dinh.json'), encoding='utf-8'))
    SD = json.load(open(os.path.join(DIR, 'ket_qua', 'on_dinh_suon_doi.json'), encoding='utf-8'))['chuong5']

    # ---- mặt cắt lòng sông: TH1 ÷ TH10 (chay_tham_on_dinh.py) + TH11, TH12 (phan_tich_bo_sung.py)
    ls_ht = {r['ma']: r for k, r in T.items() if k.startswith('HT_')}
    ls_sc = {r['ma']: dict(r) for k, r in T.items() if k.startswith('SC_')}
    for ma, khoa, mai in (('TH11', 'TC_HL_q20', 'HL'), ('TH12', 'TC_TL_q20', 'TL')):
        ls_sc[ma] = dict(ma=ma, mo_ta='Thi công: MNDBT + tải thi công %s kPa trên đỉnh' % so(B[khoa]['q'], 0), mai=mai,
                         to_hop='TC', K=B[khoa]['K_bishop'], K_cp=1.24)
    # ---- mặt cắt sườn đồi (V – T9)
    sd_ht = {r['ma']: r for r in SD['HT']}
    sd_sc = {r['ma']: r for r in SD['SC']}
    a, b = danh_gia(ls_ht, ls_sc), danh_gia(sd_ht, sd_sc)
    hon_hl = [100 * (sd_sc[m]['K'] / ls_sc[m]['K'] - 1) for m in sd_sc if sd_sc[m]['mai'] == 'HL']
    kem = [m for m in sd_sc if sd_sc[m]['K'] < ls_sc[m]['K']]
    du_tru_min = min([('lòng sông', r) for r in ls_sc.values()] + [('sườn đồi', r) for r in sd_sc.values()],
                     key=lambda t: t[1]['K'] / t[1]['K_cp'])

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.0), Cm(2.5), Cm(2.5)
    doan(d, 'KẾT QUẢ KIỂM TRA ỔN ĐỊNH MÁI ĐẬP SAU SỬA CHỮA – MẶT CẮT LÒNG SÔNG VÀ MẶT CẮT SƯỜN ĐỒI', co=13, dam=True,
         can=WD_ALIGN_PARAGRAPH.CENTER, truoc=0)
    doan(d, 'Đỉnh đập hiện trạng +%s m, sau sửa chữa +%s m; cơ thượng lưu, hạ lưu +%s m; MNDBT +%s m, MNLTK +%s m, '
            'MNLKT +%s m, MNC +%s m.'
         % (so(S.Z_DINH_DAP_HT), so(S.Z_DINH_DAP_SC), so(S.Z_CO_TL, 1), so(S.MNDBT), so(S.MNLTK), so(S.MNLKT),
            so(S.MNC)), co=11, nghieng=True, can=WD_ALIGN_PARAGRAPH.CENTER)
    doan(d, 'Các trường hợp tính toán, tổ hợp tải trọng và phương pháp như mục 2.5 (Bishop đơn giản, mặt trượt trụ tròn; '
            'áp lực nước lỗ rỗng lấy từ lời giải thấm của từng mặt cắt – Bảng 5.1, Bảng 5.2); bổ sung điều kiện thi công – '
            'sửa chữa (TCVN 8216:2018, mục 11.3.2 và Bảng 1): hồ ở MNDBT, tải trọng xe máy thi công trên toàn bề rộng đỉnh '
            'đập q = 20 kPa, [K] = 1,24 (TH11, TH12). Mặt cắt lòng sông lấy tại cọc C16 (Km0+400, gần tuyến quan trắc '
            'T7), mặt cắt sườn đồi theo mô hình mặt cắt V – T9; sau sửa chữa mặt đỉnh được nâng lên +%s m.'
         % so(S.Z_DINH_DAP_SC), truoc=6, thut=1.0)

    bang(d, 'Bảng 5.3 Kết quả kiểm tra ổn định mái đập mặt cắt lòng sông sau sửa chữa (Bishop)', TIEU_DE,
         dong_bang(ls_ht, ls_sc), cot=COT, cot_trai=(1,))
    doan(d, 'Tất cả các trường hợp đều thỏa mãn; hệ số ổn định nhỏ nhất của mái hạ lưu K = %s (%s – %s, [K] = %s), của '
            'mái thượng lưu K = %s (%s – %s, [K] = %s). So với hiện trạng, hệ số ổn định sau sửa chữa giảm %s ÷ %s%% vì '
            'phần đắp nâng đỉnh đập làm tăng tải trọng ở đầu các cung trượt sâu; mức giảm nhỏ, mọi trường hợp vẫn lớn '
            'hơn [K] ít nhất %d%%. Trong giai đoạn thi công, tải trọng xe máy trên đỉnh làm hệ số ổn định giảm %s (mái '
            'hạ lưu, TH11 so với TH1) và %s (mái thượng lưu, TH12 so với TH7), vẫn lớn hơn [K] = 1,24.'
         % (so(a['hl']['K'], 3), a['hl']['ma'], ten_th(a['hl']), so(a['hl']['K_cp']), so(a['tl']['K'], 3),
            a['tl']['ma'], ten_th(a['tl']), so(a['tl']['K_cp']), so(a['giam_min'], 1),
            so(a['giam_max'], 1), int(a['du_tru']), so(a['tc_hl']), so(a['tc_tl'])), truoc=6, thut=1.0)

    bang(d, 'Bảng 5.4 Kết quả kiểm tra ổn định mái đập mặt cắt sườn đồi sau sửa chữa (Bishop)', TIEU_DE,
         dong_bang(sd_ht, sd_sc), cot=COT, cot_trai=(1,))
    cau = ('Tất cả các trường hợp đều thỏa mãn; hệ số ổn định nhỏ nhất của mái hạ lưu K = %s (%s – %s, [K] = %s), của '
           'mái thượng lưu K = %s (%s – %s, [K] = %s). Mặt đỉnh chỉ được nâng trong phạm vi bề rộng đỉnh (+%s → +%s m, '
           'lớp đắp bù đất sét pha K ≥ 0,97); tải trọng lớp đắp làm hệ số ổn định giảm %s ÷ %s%% so với hiện trạng, tải '
           'trọng xe máy thi công làm giảm thêm %s (mái hạ lưu, TH11 so với TH1) và %s (mái thượng lưu, TH12 so với TH7), '
           'vẫn lớn hơn [K] = 1,24. So với mặt cắt lòng sông, mái hạ lưu của mặt cắt sườn đồi ổn định hơn (K lớn hơn '
           '%s ÷ %s%%) vì chiều cao đập nhỏ hơn (mặt đất hạ lưu +68,3 m so với +54,0 m ở lòng sông)'
           % (so(b['hl']['K'], 3), b['hl']['ma'], ten_th(b['hl']), so(b['hl']['K_cp']), so(b['tl']['K'], 3),
              b['tl']['ma'], ten_th(b['tl']), so(b['tl']['K_cp']), so(S.Z_DINH_DAP_HT), so(S.Z_DINH_DAP_SC),
              so(b['giam_min'], 1), so(b['giam_max'], 1), so(b['tc_hl']), so(b['tc_tl']), so(min(hon_hl), 0),
              so(max(hon_hl), 0)))
    if kem:
        cau += '; riêng mái thượng lưu %s có K nhỏ hơn mặt cắt lòng sông (%s)' % (', '.join(kem), '; '.join(
            '%s: %s so với %s' % (m, so(sd_sc[m]['K'], 3), so(ls_sc[m]['K'], 3)) for m in kem))
    if 'TH8' in kem:   # cùng hạ MNDBT → MNC nhưng đập ở sườn đồi thấp: phần nước đỡ mái mất đi lớn hơn nhiều
        ha = S.MNDBT - S.MNC
        cau += (': cùng hạ %s m nhưng ở sườn đồi mực nước hạ bằng khoảng %d%% chiều cao đập (chân mái thượng lưu +%s m), '
                'ở lòng sông chỉ khoảng %d%% (+%s m)'
                % (so(ha, 1), round(100 * ha / (S.Z_DINH_DAP_HT - Z_CHAN_SD)), so(Z_CHAN_SD, 1),
                   round(100 * ha / (S.Z_DINH_DAP_HT - Z_CHAN_LS)), so(Z_CHAN_LS, 1)))
    ten_mc, r_min = du_tru_min
    cau += ('. Trường hợp có dự trữ nhỏ nhất của toàn tuyến đập là %s – %s ở mặt cắt %s (K = %s, lớn hơn [K] = %s '
            'khoảng %s%%).' % (r_min['ma'], ten_th(r_min), ten_mc, so(r_min['K'], 3), so(r_min['K_cp']),
                               so(100 * (r_min['K'] / r_min['K_cp'] - 1), 0)))
    doan(d, cau, truoc=6, thut=1.0)
    f_kc = os.path.join(DIR, 'ket_qua', 'kiem_chung_cung_truot_sau_sc.json')    # kt_cung_truot_sau_sc.py
    if os.path.exists(f_kc):
        KC = json.load(open(f_kc, encoding='utf-8'))
        lech = {k: max(0.0, max(r['giam_pt'] for tt in KC[k].values() for r in tt.values())) for k in KC}
        doan(d, 'Kiểm tra việc dò tâm trượt: tối ưu hóa Nelder – Mead từ 10 điểm xuất phát ở các vùng tâm khác nhau cho '
                'hệ số ổn định nhỏ nhất thấp hơn giá trị trong Bảng 5.3, Bảng 5.4 không quá %s%% (lòng sông) và %s%% (sườn '
                'đồi), ở cả hiện trạng và sau sửa chữa – kết luận không đổi.'
             % (so(lech['long_song'], 1), so(lech['suon_doi'], 1)), truoc=6, thut=1.0)
    if os.path.exists(HINH_SD):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(HINH_SD, width=Cm(15.5))
        doan(d, 'Hình 5.5 Cung trượt nguy hiểm mặt cắt sườn đồi sau sửa chữa (TH1 ÷ TH12)', co=12,
             can=WD_ALIGN_PARAGRAPH.CENTER, nghieng=True, truoc=3, sau=6)
    doan(d, 'Ghi chú: HL, TL – mái hạ lưu, thượng lưu; CB, ĐB, TC – tổ hợp cơ bản, đặc biệt, thi công; [K] theo TCVN '
            '8216:2018 Bảng 1, đập cấp II. Động đất: hệ số giả tĩnh k_{h} = 0,05. Mưa kéo dài: tầng bão hòa song song mái '
            'hạ lưu dày 3 m (hiện trạng) và 1 m (sau sửa chữa, mái có cỏ và rãnh thoát nước). Hạ, rút mực nước: áp lực '
            'nước lỗ rỗng giữ như thấm ổn định ở mực nước cao (đường bão hòa chưa kịp hạ), tải trọng nước trên mái lấy '
            'theo mực nước thấp. Mặt cắt sườn đồi: mặt đất hạ lưu +68,3 m cao hơn mực nước hạ lưu nên mọi trường hợp '
            'không có nước ở hạ lưu. Thi công (TH11, TH12) chỉ xét cho mặt cắt sau sửa chữa.',
         co=11, nghieng=True, truoc=6)
    d.save(RA)
    print('Đã ghi', os.path.normpath(RA))
    for ten, ht, sc in (('Lòng sông', ls_ht, ls_sc), ('Sườn đồi', sd_ht, sd_sc)):
        print(ten)
        for dong in dong_bang(ht, sc):
            print('  ' + ' | '.join(dong))


if __name__ == '__main__':
    main()
