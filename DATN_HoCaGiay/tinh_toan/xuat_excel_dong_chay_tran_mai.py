# -*- coding: utf-8 -*-
"""
XUẤT FILE EXCEL: DÒNG CHẢY TRÀN TRÊN MÁI HẠ LƯU KHI MƯA – BẢNG 5.5 (mục 5.4.1 thuyết minh) VÀ KIỂM TRA KHẢ NĂNG THOÁT
NƯỚC CỦA RÃNH – BẢNG 5.6 (mục 5.4.2), CÓ CÔNG THỨC
  q = C·i·L ;  h = (q·n/√sinβ)^0,6 ;  τ = γw·h·sinβ ;  v = q/h          (Chow [25] – dòng màng mỏng, Manning)
Hai đoạn mái hạ lưu mặt cắt lòng sông sau sửa chữa: m = 3,0 (đỉnh +78,85 → cơ +68,0) và m = 3,5 (cơ +68,0 → đỉnh lăng
trụ +59,0); ô rãnh 4 × 4 m xoay 45° nên đường chảy dài nhất là đường chéo ô. Kèm độ nhạy (C, n, tỷ lệ mưa 30 phút)
và biểu đồ τ ~ L. Rãnh: Qtk = C·i·A, Qtn = (1/n)·ω·R^(2/3)·√S ở chiều sâu 75% chiều cao rãnh, chiều sâu và lưu tốc khi
xả Qtk (lặp), kiểm tra bất lợi. Mọi ô kết quả là công thức, ô vàng chữ xanh là số liệu nhập.
  python3 xuat_excel_dong_chay_tran_mai.py [đường dẫn .xlsx ra]
Mặc định: ../TAI_VE/DongChayTranMai_Bang5_5.xlsx
"""
import math
import os
import sys

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

DIR = os.path.dirname(os.path.abspath(__file__))
RA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DIR, '..', 'TAI_VE', 'DongChayTranMai_Bang5_5.xlsx')

FONT = 'Times New Roman'
f_thuong = Font(name=FONT, size=11)
f_dam = Font(name=FONT, size=11, bold=True)
f_tieu_de = Font(name=FONT, size=14, bold=True, color='12324A')
f_muc = Font(name=FONT, size=12, bold=True, color='A65A10')
f_nhap = Font(name=FONT, size=11, color='0000FF', bold=True)
f_trang = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_ghi_chu = Font(name=FONT, size=10, italic=True, color='52514E')
nen_nhap = PatternFill('solid', fgColor='FFF2CC')
nen_tieu_de = PatternFill('solid', fgColor='12324A')
nen_kq = PatternFill('solid', fgColor='FBF3E8')
vien = Border(*(Side(style='thin', color='B8BEC6'),) * 4)
giua = Alignment(horizontal='center', vertical='center', wrap_text=True)
trai = Alignment(horizontal='left', vertical='center', wrap_text=True)

wb = Workbook()


def ghi(ws, o, gt, font=f_thuong, fmt=None, fill=None, can=None, border=True):
    c = ws[o]
    c.value = gt
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if can:
        c.alignment = can
    if border:
        c.border = vien
    return c


def ten(name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def tieu_de_cot(ws, hang, cot0, nhan, cao=None):
    for k, t in enumerate(nhan):
        ghi(ws, '%s%d' % (get_column_letter(cot0 + k), hang), t, f_trang, fill=nen_tieu_de, can=giua)
    if cao:
        ws.row_dimensions[hang].height = cao


def do_rong(ws, rong):
    for k, w in enumerate(rong):
        ws.column_dimensions[get_column_letter(k + 1)].width = w


def in_an(ws, ngang=True, mot_trang=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if ngang else 'portrait'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if mot_trang else 0
    ws.page_margins.left = ws.page_margins.right = 0.4


# ================================================================== HƯỚNG DẪN
hd = wb.active
hd.title = 'HuongDan'
do_rong(hd, (4, 120))
ghi(hd, 'A1', 'DÒNG CHẢY TRÀN TRÊN MÁI HẠ LƯU KHI MƯA – CHIA Ô MÁI VÀ KIỂM TRA RÃNH THOÁT NƯỚC (mục 5.4.1, 5.4.2)',
    f_tieu_de, border=False)
DONG = [
    ('M', '1. Mục đích'),
    ('-', 'Mái hạ lưu bị xói rãnh do nước mưa chảy tràn; lực xói tăng theo chiều dài đường chảy L. Chia mái thành ô bằng rãnh '
          'xiên 45° (TCVN 8216:2018 điều 10.3.4.5) để cắt ngắn L; kích thước rãnh, ô xác định qua tính toán (điều 10.3.4.6).'),
    ('M', '2. Công thức (Chow, Open-Channel Hydraulics [25])'),
    ('c', '(5-1)  q = C·i·L          – lưu lượng đơn vị (m²/s): mái dài L phía trên điểm xét hứng mưa trên L × 1 m².'),
    ('c', '(5-2)  h = (q·n/√sinβ)^0,6 – chiều sâu dòng màng mỏng, công thức Manning cho dòng rất rộng (R ≈ h).'),
    ('c', '(5-3)  τ = γw·h·sinβ      – ứng suất tiếp đáy (lực kéo của dòng chảy đều), so với ứng suất tiếp cho phép của đất.'),
    ('c', '         v = q/h           – lưu tốc trung bình;  t tập trung = (n·L/√sinβ)^0,6 / (C·i)^0,4 – sóng động học.'),
    ('-', 'Gộp lại: τ = γw·(C·i·n·L)^0,6·(sinβ)^0,7 → τ tỷ lệ với L^0,6: mức giảm τ khi rút ngắn L KHÔNG phụ thuộc C, i, n '
          '(xem sheet DoNhay).'),
    ('M', '3. Ý nghĩa các chiều dài đường chảy L (sheet Bang5_5)'),
    ('-', 'Cả đoạn mái (34 m với m = 3; 32,8 m với m = 3,5): chưa có rãnh – hiện trạng (rãnh cũ hư hỏng).'),
    ('-', 'Đường chéo ô 4 × 4 m = 5,66 m: đường chảy tính toán sau sửa chữa – các ô đặt xoay 45° so với đường dốc nhất nên '
          'nước chảy theo đường chéo ô. Cạnh ô 4 m chỉ để tham khảo.'),
    ('-', '17 m, 8 m: các bước chia đôi đường chảy (thêm 1 hoặc 3 rãnh ngang) – để thấy τ giảm dần theo L.'),
    ('-', 'Đoạn mái trên (m = 3) là đoạn quyết định: dốc hơn và dài hơn; rãnh ngang trên cơ +68,0 thu nước đoạn trên nên '
          'đoạn dưới (m = 3,5) chỉ chịu nước mưa rơi trên chính nó.'),
    ('M', '4. Cách dùng'),
    ('-', 'Chỉ sửa các ô nền vàng chữ xanh ở sheet SoLieu (và các ô vàng ở sheet DoNhay); mọi ô khác là công thức.'),
    ('-', 'Số liệu: lượng mưa 1 ngày lớn nhất P = 10% = 191,8 mm – QTVH (2026) Bảng 2-11; mưa 30 phút = 30% mưa ngày '
          '(giả thiết – hồ sơ không có mưa thời đoạn ngắn); C = 0,6 (mặt đất dốc lớn, cỏ thưa); n = 0,03 (mặt đất trống); '
          'ứng suất tiếp cho phép của đất không có lớp phủ theo Chow [25].'),
    ('-', 'Chiều dài cả mái trên tính đúng là 34,3 m; bảng làm tròn 34 m như thuyết minh (τ chênh 0,04 Pa).'),
    ('M', '5. Kiểm tra khả năng thoát nước của rãnh (sheet Bang5_6_Ranh)'),
    ('c', 'Qtk = C·i·A      – lưu lượng thiết kế: phương pháp cường độ giới hạn, A – diện tích hứng nước (hình chiếu ngang).'),
    ('c', 'Qtn = (1/n)·ω·R^(2/3)·√S – khả năng thoát (Manning), nước sâu h = 0,75·H; ω = b·h; χ = b + 2h; R = ω/χ; v = Qtn/ω.'),
    ('-', 'Diện tích hứng nước: rãnh dọc – khoảng cách rãnh 40 m × (bề rộng đỉnh + hình chiếu ngang mái trên); rãnh ngang '
          'trên cơ – 40 m × (bề rộng cơ + dải chân mái trên); rãnh xiên – một ô 4 × 4 m; rãnh ngang đỉnh – 40 m × bề rộng đỉnh.'),
    ('-', 'Độ dốc S: rãnh dọc chạy theo đường dốc nhất S = sinβ; rãnh xiên lệch 45° trên mặt mái S = sinβ/√2; rãnh ngang '
          '(đỉnh, cơ) S = 0,003.'),
    ('-', 'Bảng b) tính chiều sâu và lưu tốc khi rãnh xả đúng Qtk (lặp 5 lần công thức Manning); bảng c) kiểm tra khi nước '
          'dồn nhiều hơn giả thiết – vẫn đạt.'),
]
r = 3
for kieu, t in DONG:
    if kieu == 'M':
        r += 1
        ghi(hd, 'A%d' % r, t, f_muc, border=False)
    else:
        ghi(hd, 'A%d' % r, '•' if kieu == '-' else '', f_thuong, can=Alignment(horizontal='center', vertical='top'),
            border=False)
        ghi(hd, 'B%d' % r, t, Font(name='Cambria Math', size=11) if kieu == 'c' else f_thuong,
            can=Alignment(horizontal='left', vertical='top', wrap_text=True), border=False)
        hd.row_dimensions[r].height = 30 if len(t) > 115 else 16
    r += 1
in_an(hd, ngang=False)

# ================================================================== SỐ LIỆU
sl = wb.create_sheet('SoLieu')
do_rong(sl, (46, 12, 9, 70))
ghi(sl, 'A1', 'SỐ LIỆU ĐẦU VÀO (ô vàng chữ xanh: nhập; còn lại: công thức)', f_tieu_de, border=False)
THONG_SO = [  # (nhãn, giá trị hoặc công thức, đơn vị, tên, ghi chú, nhập?, định dạng)
    ('A. MƯA TÍNH TOÁN', None),
    ('Lượng mưa 1 ngày lớn nhất P = 10%', 191.8, 'mm', 'Mua_ngay', 'QTVH (2026) Bảng 2-11 – thuyết minh dẫn [1]', True, '0.0'),
    ('Tỷ lệ mưa 30 phút / mưa ngày', 0.30, '–', 'Ty_le_mua', 'Giả thiết (hồ sơ không có mưa thời đoạn ngắn); độ nhạy 0,20 ÷ 0,40',
     True, '0.00'),
    ('Thời đoạn mưa', 0.5, 'giờ', 'T_mua', 'Thời gian tập trung nước trên mái chỉ 0,5 ÷ 2 phút (sheet Bang5_5)', True, '0.0'),
    ('Cường độ mưa tính toán i', '=Mua_ngay*Ty_le_mua/T_mua', 'mm/h', 'i_mmh', 'i = 191,8 × 0,30 / 0,5', False, '0.0'),
    ('Cường độ mưa tính toán i', '=i_mmh/3600000', 'm/s', 'i_ms', 'đổi đơn vị: mm/h → m/s', False, '0.000E+00'),
    ('B. MẶT MÁI', None),
    ('Hệ số dòng chảy C', 0.6, '–', 'HS_C', 'Mặt đất dốc lớn, cỏ thưa khi mưa lớn (bảng hệ số dòng chảy ≈ 0,4 ÷ 0,6), lấy cận '
     'trên', True, '0.00'),
    ('Hệ số nhám Manning n', 0.03, '–', 'HS_n', 'Mặt đất trống – Chow [25]; dòng màng mỏng trên đất trống ≈ 0,011 ÷ 0,05', True,
     '0.000'),
    ('Trọng lượng riêng của nước γw', 9810, 'N/m³', 'Gamma_w', '', True, '0'),
    ('C. HÌNH HỌC MÁI HẠ LƯU (mặt cắt lòng sông sau sửa chữa)', None),
    ('Cao trình đỉnh đập', 78.85, 'm', 'Z_dinh', 'Sau sửa chữa', True, '0.00'),
    ('Cao trình cơ hạ lưu', 68.0, 'm', 'Z_co', 'Có rãnh ngang 40 × 40 cm thu nước đoạn mái trên', True, '0.00'),
    ('Cao trình đỉnh lăng trụ thoát nước', 59.0, 'm', 'Z_ltr', '', True, '0.00'),
    ('Hệ số mái đoạn trên m₁', 3.0, '–', 'm_tren', 'Từ đỉnh xuống cơ', True, '0.0'),
    ('Hệ số mái đoạn dưới m₂', 3.5, '–', 'm_duoi', 'Từ cơ xuống đỉnh lăng trụ', True, '0.0'),
    ('sinβ₁ = 1/√(1 + m₁²)', '=1/SQRT(1+m_tren^2)', '–', 'sin_tren', 'β₁ = 18,43°', False, '0.0000'),
    ('sinβ₂ = 1/√(1 + m₂²)', '=1/SQRT(1+m_duoi^2)', '–', 'sin_duoi', 'β₂ = 15,95°', False, '0.0000'),
    ('Chiều dài đoạn mái trên (dọc mái)', '=(Z_dinh-Z_co)*SQRT(1+m_tren^2)', 'm', 'L_tren', 'Đường chảy khi chưa có rãnh',
     False, '0.00'),
    ('Chiều dài đoạn mái dưới (dọc mái)', '=(Z_co-Z_ltr)*SQRT(1+m_duoi^2)', 'm', 'L_duoi', '', False, '0.00'),
    ('Cạnh ô rãnh a', 4.0, 'm', 'a_o', 'Ô 4 × 4 m (TKCS, Hình 3.2), đặt xoay 45°', True, '0.00'),
    ('Đường chéo ô a·√2', '=a_o*SQRT(2)', 'm', 'L_cheo', 'Đường chảy dài nhất trong ô (theo đường dốc nhất)', False, '0.00'),
    ('D. ỨNG SUẤT TIẾP CHO PHÉP CỦA ĐẤT KHÔNG CÓ LỚP PHỦ (Chow [25])', None),
    ('Đất pha cát', 1.8, 'Pa', 'tau_cat', '', True, '0.0'),
    ('Đất bụi', 2.3, 'Pa', 'tau_bui', '', True, '0.0'),
    ('Đất thịt chặt', 3.6, 'Pa', 'tau_thit', 'Khi cỏ đã phủ kín, mái chịu được ứng suất tiếp lớn hơn nhiều lần', True, '0.0'),
    ('E. RÃNH THOÁT NƯỚC MÁI (Bảng 5.6)', None),
    ('Hệ số nhám rãnh bê tông n', 0.015, '–', 'n_BT', 'Bê tông', True, '0.000'),
    ('Chiều sâu nước tính khả năng thoát / chiều cao rãnh', 0.75, '–', 'Ty_le_sau', 'Chừa 25% chiều cao an toàn', True, '0.00'),
    ('Khoảng cách rãnh dọc', 40.0, 'm', 'K_doc', 'TCVN 8216:2018 điều 10.3.4.6: 25 ÷ 50 m', True, '0.0'),
    ('Bề rộng đỉnh đập', 5.0, 'm', 'B_dinh', 'Dốc ngang về hạ lưu, nước qua khe gờ chắn bánh vào rãnh ngang đỉnh', True, '0.0'),
    ('Bề rộng cơ hạ lưu', 14.0, 'm', 'B_co', '', True, '0.0'),
    ('Dải chân mái trên chảy vào rãnh cơ', 3.0, 'm', 'D_chan', 'Giả thiết (khớp diện tích 680 m² của Bảng 5.6); xem bảng c) '
     'kiểm tra bất lợi', True, '0.0'),
    ('Độ dốc dọc rãnh ngang (đỉnh, cơ)', 0.003, '–', 'i_ngang', '', True, '0.000'),
    ('Hình chiếu ngang đoạn mái trên', '=m_tren*(Z_dinh-Z_co)', 'm', 'X_tren', '', False, '0.00'),
    ('Hình chiếu ngang đoạn mái dưới', '=m_duoi*(Z_co-Z_ltr)', 'm', 'X_duoi', '', False, '0.00'),
]
r = 3
for dong in THONG_SO:
    if dong[1] is None:
        r += 1
        ghi(sl, 'A%d' % r, dong[0], f_muc, border=False)
        r += 1
        tieu_de_cot(sl, r, 1, ['Thông số', 'Giá trị', 'Đơn vị', 'Ghi chú, nguồn'])
        r += 1
        continue
    nhan, gt, dv, name, gc, nhap, fmt = dong
    ghi(sl, 'A%d' % r, nhan, f_dam)
    ghi(sl, 'B%d' % r, gt, f_nhap if nhap else f_dam, fmt, nen_nhap if nhap else nen_kq, giua)
    ghi(sl, 'C%d' % r, dv, can=giua)
    ghi(sl, 'D%d' % r, gc, f_ghi_chu, can=trai)
    ten(name, 'SoLieu!$B$%d' % r)
    r += 1
in_an(sl, ngang=False)

# ================================================================== BẢNG 5.5
bg = wb.create_sheet('Bang5_5')
do_rong(bg, (10, 34, 13, 10, 10, 10, 13, 30, 12))
ghi(bg, 'A1', 'BẢNG 5.5 – ĐẶC TRƯNG DÒNG CHẢY TRÀN TRÊN MÁI HẠ LƯU THEO CHIỀU DÀI ĐƯỜNG CHẢY', f_tieu_de, border=False)
ghi(bg, 'A2', 'q = C·i·L;  h = (q·n/√sinβ)^0,6;  v = q/h;  τ = γw·h·sinβ. Ô vàng chữ xanh (cột L) nhập được; dòng in đậm '
              'dùng để kết luận.', f_ghi_chu, border=False)
COT = ['L (m)', 'Ý nghĩa', 'q = C·i·L (m²/s)', 'h (mm)', 'v (m/s)', 'τ (Pa)', 'Mức giảm so với cả đoạn mái (lần)',
       'So với ứng suất tiếp cho phép', 'Thời gian tập trung (s)']


def bang_doan(r0, tieu_de, sin_ten, dong_l):
    """dong_l: (L – số hoặc công thức, ý nghĩa, in đậm?); dòng cuối là cả đoạn mái (mốc so sánh)."""
    ghi(bg, 'A%d' % r0, tieu_de, f_muc, border=False)
    tieu_de_cot(bg, r0 + 1, 1, COT, cao=45)
    r_ca = r0 + 1 + len(dong_l)
    for k, (L, y_nghia, dam) in enumerate(dong_l):
        r = r0 + 2 + k
        f = f_dam if dam else f_thuong
        nhap = not (isinstance(L, str) and L.startswith('='))
        fmt_L = '0.00' if L == '=L_cheo' else 'General'
        ghi(bg, 'A%d' % r, L, f_nhap if nhap else f, fmt_L, nen_nhap if nhap else None, giua)
        ghi(bg, 'B%d' % r, y_nghia, f, can=trai)
        ghi(bg, 'C%d' % r, '=HS_C*i_ms*A%d' % r, f, '0.00E+00', can=giua)
        ghi(bg, 'D%d' % r, '=1000*(C%d*HS_n/SQRT(%s))^0.6' % (r, sin_ten), f, '0.00', can=giua)
        ghi(bg, 'E%d' % r, '=C%d/(D%d/1000)' % (r, r), f, '0.00', can=giua)
        ghi(bg, 'F%d' % r, '=Gamma_w*D%d/1000*%s' % (r, sin_ten), f, '0.00', nen_kq, giua)
        ghi(bg, 'G%d' % r, '=F$%d/F%d' % (r_ca, r), f, '0.0', can=giua)
        ghi(bg, 'H%d' % r, '=IF(F{r}<=tau_cat,"Nhỏ hơn [τ] đất pha cát",IF(F{r}<=tau_bui,"Nhỏ hơn [τ] đất bụi",'
                           'IF(F{r}<=tau_thit,"Nhỏ hơn [τ] đất thịt chặt","Lớn hơn [τ] mọi loại đất trống")))'.format(r=r),
            f, can=trai)
        ghi(bg, 'I%d' % r, '=(HS_n*A{r}/SQRT({s}))^0.6/(HS_C*i_ms)^0.4'.format(r=r, s=sin_ten), f, '0', can=giua)
    return r_ca


R1 = 4
r_ca1 = bang_doan(R1, 'a) Đoạn mái trên m = 3,0 (đỉnh +78,85 → cơ +68,0) – đoạn quyết định, dùng cho Bảng 5.5', 'sin_tren', [
    ('=a_o', 'Cạnh ô 4 × 4 m (tham khảo)', False),
    ('=L_cheo', 'Đường chéo ô – đường chảy tính toán sau sửa chữa', True),
    (8, 'Thêm 3 rãnh ngang (so sánh)', False),
    (17, 'Thêm 1 rãnh ngang giữa mái (so sánh)', False),
    ('=ROUND(L_tren,0)', 'Cả đoạn mái, chưa có rãnh – hiện trạng', True)])
R2 = r_ca1 + 3
r_ca2 = bang_doan(R2, 'b) Đoạn mái dưới m = 3,5 (cơ +68,0 → đỉnh lăng trụ +59,0) – kiểm tra thêm', 'sin_duoi', [
    ('=a_o', 'Cạnh ô 4 × 4 m (tham khảo)', False),
    ('=L_cheo', 'Đường chéo ô – đường chảy tính toán sau sửa chữa', True),
    (8, 'Thêm 3 rãnh ngang (so sánh)', False),
    (17, 'Thêm 1 rãnh ngang giữa mái (so sánh)', False),
    ('=ROUND(L_duoi,1)', 'Cả đoạn mái, chưa có rãnh', True)])
R3 = r_ca2 + 3
ghi(bg, 'A%d' % R3, 'c) Tổng hợp mức giảm ứng suất tiếp nhờ chia ô 4 × 4 m', f_muc, border=False)
tieu_de_cot(bg, R3 + 1, 1, ['', 'Nội dung', 'Đoạn trên m = 3,0', 'Đoạn dưới m = 3,5'], cao=30)
TONG = [
    ('τ cả đoạn mái, chưa có rãnh (Pa)', '=F%d' % r_ca1, '=F%d' % r_ca2, '0.00'),
    ('τ trên ô, theo đường chéo (Pa)', '=F%d' % (R1 + 3), '=F%d' % (R2 + 3), '0.00'),
    ('Mức giảm theo đường chéo ô (lần)', '=G%d' % (R1 + 3), '=G%d' % (R2 + 3), '0.0'),
    ('Mức giảm theo cạnh ô (lần)', '=G%d' % (R1 + 2), '=G%d' % (R2 + 2), '0.0'),
    ('Kiểm tra: (L cả mái / L chéo)^0,6 – không phụ thuộc C, i, n', '=(A%d/A%d)^0.6' % (r_ca1, R1 + 3),
     '=(A%d/A%d)^0.6' % (r_ca2, R2 + 3), '0.00'),
]
for k, (nhan, a, b, fmt) in enumerate(TONG):
    r = R3 + 2 + k
    ghi(bg, 'A%d' % r, k + 1, can=giua)
    ghi(bg, 'B%d' % r, nhan, can=trai)
    ghi(bg, 'C%d' % r, a, f_dam, fmt, nen_kq, giua)
    ghi(bg, 'D%d' % r, b, f_dam, fmt, nen_kq, giua)
R_GC = R3 + 2 + len(TONG) + 1
for k, t in enumerate([
        'Ghi chú: cột "Mức giảm" = τ cả đoạn mái / τ của dòng đó. Thời gian tập trung nước trên mái chỉ 0,5 ÷ 2 phút nên '
        'dòng chảy ổn định ngay trong trận mưa – công thức (5-1) áp dụng được.',
        'Rãnh ngang trên cơ +68,0 phải được duy tu thông thoáng: nếu bị lấp, nước của đoạn mái trên và mặt cơ tràn xuống '
        'đoạn dưới (đường chảy ≈ 81 m) làm τ ở chân mái dưới tăng tới khoảng 10 Pa.']):
    ghi(bg, 'A%d' % (R_GC + k), t, f_ghi_chu, border=False)
bg.freeze_panes = 'A4'
in_an(bg, mot_trang=True)

# ================================================================== ĐỘ NHẠY
dn = wb.create_sheet('DoNhay')
do_rong(dn, (32, 9, 12, 9, 11, 15, 15, 13))
ghi(dn, 'A1', 'ĐỘ NHẠY THEO CÁC THÔNG SỐ GIẢ THIẾT (đoạn mái trên m = 3,0)', f_tieu_de, border=False)
ghi(dn, 'A2', 'Ô vàng: thông số của từng trường hợp (sửa được). τ = γw·(C·i·n·L/√sinβ)^0,6·sinβ. Mức giảm không đổi khi C, '
              'i, n thay đổi.', f_ghi_chu, border=False)
tieu_de_cot(dn, 4, 1, ['Trường hợp', 'C', 'Tỷ lệ mưa 30 phút', 'n', 'i (mm/h)', 'τ trên ô – đường chéo (Pa)',
                       'τ cả đoạn mái (Pa)', 'Mức giảm (lần)'], cao=45)
TH = [('Cơ sở (như sheet SoLieu)', '=HS_C', '=Ty_le_mua', '=HS_n'),
      ('Mưa 30 phút = 20% mưa ngày', '=HS_C', 0.20, '=HS_n'),
      ('Mưa 30 phút = 40% mưa ngày', '=HS_C', 0.40, '=HS_n'),
      ('Đất gần như không thấm, C = 0,9', 0.9, '=Ty_le_mua', '=HS_n'),
      ('Mặt đất trơn, n = 0,011', '=HS_C', '=Ty_le_mua', 0.011),
      ('Mặt đất gồ ghề, n = 0,05', '=HS_C', '=Ty_le_mua', 0.05)]
for k, (nhan, C, ty, n) in enumerate(TH):
    r = 5 + k
    ghi(dn, 'A%d' % r, nhan, f_dam if k == 0 else f_thuong, can=trai)
    for col, gt, fmt in (('B', C, '0.00'), ('C', ty, '0.00'), ('D', n, '0.000')):
        nhap = not (isinstance(gt, str) and gt.startswith('='))
        ghi(dn, '%s%d' % (col, r), gt, f_nhap if nhap else f_thuong, fmt, nen_nhap if nhap else None, giua)
    ghi(dn, 'E%d' % r, '=Mua_ngay*C%d/T_mua' % r, fmt='0.0', can=giua)
    for col, L in (('F', 'Bang5_5!$A$%d' % (R1 + 3)), ('G', 'Bang5_5!$A$%d' % r_ca1)):
        ghi(dn, '%s%d' % (col, r), '=Gamma_w*(B{r}*E{r}/3600000*{L}*D{r}/SQRT(sin_tren))^0.6*sin_tren'.format(r=r, L=L),
            f_dam, '0.00', nen_kq, giua)
    ghi(dn, 'H%d' % r, '=G%d/F%d' % (r, r), fmt='0.00', can=giua)
in_an(dn)

# ================================================================== BIỂU ĐỒ τ ~ L
bd = wb.create_sheet('BieuDo')
do_rong(bd, (9, 13, 13, 13, 13, 13))
ghi(bd, 'A1', 'ỨNG SUẤT TIẾP THEO CHIỀU DÀI ĐƯỜNG CHẢY (số liệu vẽ biểu đồ – công thức)', f_tieu_de, border=False)
tieu_de_cot(bd, 3, 1, ['L (m)', 'τ mái m = 3,0 (Pa)', 'τ mái m = 3,5 (Pa)', '[τ] đất pha cát', '[τ] đất bụi',
                       '[τ] đất thịt chặt'], cao=45)
L_BD = list(range(1, 36))
for k, L in enumerate(L_BD):
    r = 4 + k
    ghi(bd, 'A%d' % r, L, fmt='0', can=giua)
    for col, s in (('B', 'sin_tren'), ('C', 'sin_duoi')):
        ghi(bd, '%s%d' % (col, r), '=Gamma_w*(HS_C*i_ms*A{r}*HS_n/SQRT({s}))^0.6*{s}'.format(r=r, s=s), fmt='0.00',
            can=giua)
    for col, tn in (('D', 'tau_cat'), ('E', 'tau_bui'), ('F', 'tau_thit')):
        ghi(bd, '%s%d' % (col, r), '=%s' % tn, fmt='0.0', can=giua)
r_cuoi = 3 + len(L_BD)
ch = ScatterChart()
ch.title = 'Ứng suất tiếp do dòng chảy tràn trên mái hạ lưu'
ch.style = 13
ch.x_axis.title = 'Chiều dài đường chảy L (m)'
ch.y_axis.title = 'τ (Pa)'
ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = 0, 35, 5
ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = 0, 8, 1
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.legend.position = 'b'
ch.height, ch.width = 10, 18
x = Reference(bd, min_col=1, min_row=4, max_row=r_cuoi)
for col, mau, net in ((2, '2A78D6', None), (3, '1BAF7A', None), (4, 'EB6834', 'dash'), (5, 'A65A10', 'dash'),
                      (6, '52514E', 'dash')):
    s = Series(Reference(bd, min_col=col, min_row=3, max_row=r_cuoi), x, title_from_data=True)
    s.smooth = True
    s.marker.symbol = 'none'
    s.graphicalProperties.line.solidFill = mau
    s.graphicalProperties.line.width = 22225
    if net:
        s.graphicalProperties.line.dashStyle = net
    ch.series.append(s)
bd.add_chart(ch, 'H3')
in_an(bd)

# ================================================================== BẢNG 5.6 – RÃNH THOÁT NƯỚC
rn = wb.create_sheet('Bang5_6_Ranh')
do_rong(rn, (34, 8, 8, 10, 34, 11, 10, 9, 9, 9, 9, 10, 9, 10, 10))
ghi(rn, 'A1', 'BẢNG 5.6 – KIỂM TRA KHẢ NĂNG THOÁT NƯỚC CÁC LOẠI RÃNH MÁI HẠ LƯU', f_tieu_de, border=False)
ghi(rn, 'A2', 'Qtk = C·i·A;  Qtn = (1/n)·ω·R^(2/3)·√S với chiều sâu nước h = 0,75·H;  ω = b·h;  χ = b + 2h;  R = ω/χ;  '
              'v = Qtn/ω. Số liệu ở sheet SoLieu (mục A, B, C, E).', f_ghi_chu, border=False)
ghi(rn, 'A4', 'a) Khả năng thoát nước (cột in đậm là các cột của Bảng 5.6)', f_muc, border=False)
tieu_de_cot(rn, 5, 1, ['Loại rãnh', 'b (m)', 'H (m)', 'A (m²)', 'Cách lấy diện tích hứng nước A', 'Qtk = C·i·A (m³/s)',
                       'Độ dốc S', 'h = 0,75·H (m)', 'ω = b·h (m²)', 'χ = b + 2h (m)', 'R = ω/χ (m)', 'Qtn (m³/s)',
                       'v = Qtn/ω (m/s)', 'Qtn / Qtk (lần)', 'Đánh giá'], cao=45)
RANH = [  # tên, b, H, A, cách lấy A, S
    ('Rãnh dọc 40 × 40 cm (theo mái, @ 40 m)', 0.4, 0.4, '=K_doc*(B_dinh+X_tren)',
     '40 m × (đỉnh 5 m + hình chiếu mái trên)', '=sin_tren'),
    ('Rãnh ngang trên cơ 40 × 40 cm', 0.4, 0.4, '=K_doc*(B_co+D_chan)', '40 m × (cơ 14 m + dải chân mái 3 m)', '=i_ngang'),
    ('Rãnh xiên 20 × 20 cm (ô 4 × 4 m)', 0.2, 0.2, '=a_o^2', 'Một ô 4 × 4 m', '=sin_tren/SQRT(2)'),
    ('Rãnh ngang đỉnh 20 × 20 cm', 0.2, 0.2, '=K_doc*B_dinh', '40 m × đỉnh 5 m', '=i_ngang'),
]
RA0 = 6
for k, (tenr, b, H, A, cach, S) in enumerate(RANH):
    r = RA0 + k
    ghi(rn, 'A%d' % r, tenr, f_dam, can=trai)
    ghi(rn, 'B%d' % r, b, f_nhap, '0.00', nen_nhap, giua)
    ghi(rn, 'C%d' % r, H, f_nhap, '0.00', nen_nhap, giua)
    ghi(rn, 'D%d' % r, A, f_dam, '#,##0', nen_kq, giua)
    ghi(rn, 'E%d' % r, cach, f_ghi_chu, can=trai)
    ghi(rn, 'F%d' % r, '=HS_C*i_ms*D%d' % r, f_dam, '0.0000', nen_kq, giua)
    ghi(rn, 'G%d' % r, S, f_dam, '0.000', nen_kq, giua)
    ghi(rn, 'H%d' % r, '=Ty_le_sau*C%d' % r, fmt='0.000', can=giua)
    ghi(rn, 'I%d' % r, '=B{r}*H{r}'.format(r=r), fmt='0.0000', can=giua)
    ghi(rn, 'J%d' % r, '=B{r}+2*H{r}'.format(r=r), fmt='0.000', can=giua)
    ghi(rn, 'K%d' % r, '=I{r}/J{r}'.format(r=r), fmt='0.0000', can=giua)
    ghi(rn, 'L%d' % r, '=I{r}*K{r}^(2/3)*SQRT(G{r})/n_BT'.format(r=r), f_dam, '0.000', nen_kq, giua)
    ghi(rn, 'M%d' % r, '=L{r}/I{r}'.format(r=r), f_dam, '0.00', nen_kq, giua)
    ghi(rn, 'N%d' % r, '=L{r}/F{r}'.format(r=r), fmt='0', can=giua)
    ghi(rn, 'O%d' % r, '=IF(L{r}>=F{r},"Đạt","Không đạt")'.format(r=r), f_dam, can=giua)

RB0 = RA0 + len(RANH) + 2
ghi(rn, 'A%d' % RB0, 'b) Chiều sâu và lưu tốc khi rãnh xả đúng Qtk – giải lặp công thức Manning: '
                     'h = (Qtk·n/√S)^0,6·(b + 2h)^0,4 / b', f_muc, border=False)
tieu_de_cot(rn, RB0 + 1, 1, ['Loại rãnh', 'Qtk (m³/s)', 'h₀ (m)', 'h₁ (m)', 'h₂ (m)', 'h₃ (m)', 'h₄ (m)', 'h (cm)',
                             'v (m/s)', 'h / H (%)'], cao=45)
for k in range(len(RANH)):
    r, ra = RB0 + 2 + k, RA0 + k
    ghi(rn, 'A%d' % r, '=A%d' % ra, f_dam, can=trai)
    ghi(rn, 'B%d' % r, '=F%d' % ra, fmt='0.0000', can=giua)
    ghi(rn, 'C%d' % r, '=(B{r}*n_BT/(B{a}*SQRT(G{a})))^0.6'.format(r=r, a=ra), fmt='0.0000', can=giua)
    for col, truoc in (('D', 'C'), ('E', 'D'), ('F', 'E'), ('G', 'F')):
        ghi(rn, '%s%d' % (col, r), '=(B{r}*n_BT/SQRT(G{a}))^0.6*(B{a}+2*{t}{r})^0.4/B{a}'.format(r=r, a=ra, t=truoc),
            fmt='0.0000', can=giua)
    ghi(rn, 'H%d' % r, '=100*G%d' % r, f_dam, '0.0', nen_kq, giua)
    ghi(rn, 'I%d' % r, '=B{r}/(B{a}*G{r})'.format(r=r, a=ra), f_dam, '0.00', nen_kq, giua)
    ghi(rn, 'J%d' % r, '=100*G{r}/C{a}'.format(r=r, a=ra), fmt='0', can=giua)

RC0 = RB0 + 2 + len(RANH) + 1
ghi(rn, 'A%d' % RC0, 'c) Kiểm tra bất lợi – nước dồn nhiều hơn giả thiết diện tích hứng nước ở bảng a)', f_muc, border=False)
tieu_de_cot(rn, RC0 + 1, 1, ['Trường hợp', 'b (m)', 'H (m)', 'A (m²)', 'Cách lấy diện tích hứng nước A',
                             'Qtk (m³/s)', 'Độ dốc S', 'Qtn (m³/s)', 'Qtn / Qtk (lần)', 'Đánh giá'], cao=45)
BAT_LOI = [
    ('Rãnh dọc đoạn dưới cơ (m = 3,5)', 0.4, 0.4, '=K_doc*(B_dinh+X_tren+B_co+X_duoi)',
     'Nhận nước đỉnh, mái trên, cơ và mái dưới', '=sin_duoi'),
    ('Rãnh ngang trên cơ nhận cả nước mái trên', 0.4, 0.4, '=K_doc*(B_dinh+X_tren+B_co)',
     'Như rãnh dọc + cả mặt cơ', '=i_ngang'),
    ('Rãnh xiên nhận nửa khoang mái trên', 0.2, 0.2, '=K_doc/2*X_tren', '20 m × hình chiếu mái trên',
     '=sin_tren/SQRT(2)'),
]
for k, (tenr, b, H, A, cach, S) in enumerate(BAT_LOI):
    r = RC0 + 2 + k
    ghi(rn, 'A%d' % r, tenr, f_dam, can=trai)
    ghi(rn, 'B%d' % r, b, f_nhap, '0.00', nen_nhap, giua)
    ghi(rn, 'C%d' % r, H, f_nhap, '0.00', nen_nhap, giua)
    ghi(rn, 'D%d' % r, A, f_dam, '#,##0', nen_kq, giua)
    ghi(rn, 'E%d' % r, cach, f_ghi_chu, can=trai)
    ghi(rn, 'F%d' % r, '=HS_C*i_ms*D%d' % r, f_dam, '0.0000', nen_kq, giua)
    ghi(rn, 'G%d' % r, S, fmt='0.000', can=giua)
    ghi(rn, 'H%d' % r, '=(B{r}*Ty_le_sau*C{r})*((B{r}*Ty_le_sau*C{r})/(B{r}+2*Ty_le_sau*C{r}))^(2/3)*SQRT(G{r})/n_BT'
        .format(r=r), f_dam, '0.000', nen_kq, giua)
    ghi(rn, 'I%d' % r, '=H{r}/F{r}'.format(r=r), fmt='0.0', can=giua)
    ghi(rn, 'J%d' % r, '=IF(H{r}>=F{r},"Đạt","Không đạt")'.format(r=r), f_dam, can=giua)
R_GC2 = RC0 + 2 + len(BAT_LOI) + 1
for k, t in enumerate([
        'Ghi chú: v ở bảng a) là lưu tốc khi nước sâu 75% chiều cao rãnh; khi xả đúng Qtk (bảng b) rãnh dọc chỉ ngập khoảng '
        '2,5 cm nhưng lưu tốc ≈ 2,9 m/s – rãnh phải bằng bê tông liền khối, có gờ tiêu năng ở chỗ nối với cơ.',
        'So với Bảng 5.6 trong thuyết minh: A rãnh dọc 1.460 m² = 40 × (5 + 31,5) ứng với đỉnh +78,50; với đỉnh sau sửa chữa '
        '+78,85 là 1.502 m² (Qtk 0,0288 m³/s). Rãnh xiên: bảng cũ lấy độ dốc ≈ 0,24 (= tanβ/√2); theo phương dọc rãnh là '
        'sinβ/√2 = 0,224 nên Qtn 0,145 m³/s, v 4,83 m/s (bảng cũ 0,149; 4,96). Kết luận không đổi.']):
    ghi(rn, 'A%d' % (R_GC2 + k), t, f_ghi_chu, can=Alignment(horizontal='left', vertical='top', wrap_text=True),
        border=False)
    rn.merge_cells('A%d:O%d' % (R_GC2 + k, R_GC2 + k))
    rn.row_dimensions[R_GC2 + k].height = 30
rn.freeze_panes = 'B6'
in_an(rn, mot_trang=True)

for ws, mau in ((hd, 'A65A10'), (sl, 'EDA100'), (bg, '1BAF7A'), (dn, '2A78D6'), (bd, '8A8984'), (rn, 'EB6834')):
    ws.sheet_properties.tabColor = mau
for ws in wb.worksheets:                       # mở file ở sheet Bang5_5, không chọn nhóm nhiều sheet
    ws.sheet_view.tabSelected = ws is bg
wb.active = bg
wb.save(RA)
print('Đã ghi', os.path.normpath(RA))

# kiểm tra nhanh bằng Python (cùng công thức) – để đối chiếu với giá trị Excel tính ra
i = 191.8 * 0.30 / 0.5 / 3600000
for m, Ls in ((3.0, (4, 4 * math.sqrt(2), 8, 17, 34)), (3.5, (4, 4 * math.sqrt(2), 8, 17, 32.8))):
    s = 1 / math.sqrt(1 + m * m)
    print('m = %.1f:' % m, ', '.join('L %.2f → τ %.2f Pa' % (L, 9810 * (0.6 * i * L * 0.03 / math.sqrt(s)) ** 0.6 * s)
                                     for L in Ls))
